"""Calendario sheet: twelve monthly mini-calendars, read only.

Every cell is a formula pointing at that day's row in the Presenze sheet.
Nothing is typed here: shifts are entered in Presenze.

Sheet name, headings and legend text stay in Italian: Vanessa reads them.
"""

import calendar
from datetime import date

from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from common import (BOX, NAVY, SHIFTS, MONTHS, italian_holidays, rota_row,
                    time_range)

SHORT_DAYS = ["Lun", "Mar", "Mer", "Gio", "Ven", "Sab", "Dom"]

HOURS_COL = 8        # column H, the week's hour total
FIRST_BLOCK_ROW = 10  # the legend sits above

# Colour of the row holding the day number.
WEEKDAY_FILL = "FFFFFF"
SATURDAY_FILL = "DCE9F5"
SUNDAY_FILL = "F3DCE4"
HOLIDAY_FILL = "FAE6B8"
OUTSIDE_FILL = "F7F9FB"  # days of the previous or next month

# Colour of the code and time rows, applied with conditional formatting
# because the code is typed after the file has been generated.
SHIFT_FILL = {
    "L":  "EDF1F5",
    "M":  "FFF4CC",
    "M1": "FFE8A3",
    "P":  "FFE0C2",
    "P1": "FFCFA0",
}


def _day_fill(d, holidays):
    """Holiday beats Sunday, Sunday beats Saturday."""
    if d in holidays:
        return HOLIDAY_FILL
    if d.weekday() == 6:
        return SUNDAY_FILL
    if d.weekday() == 5:
        return SATURDAY_FILL
    return WEEKDAY_FILL


def _block_header(ws, row):
    """Dark row Lun..Dom + Ore."""
    for i, v in enumerate(SHORT_DAYS + ["Ore"], start=1):
        c = ws.cell(row=row, column=i, value=v)
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BOX
    ws.row_dimensions[row].height = 20


def _month_block(ws, row, year, month, lookup, holidays, areas):
    """Writes one month's block. Returns the first free row after it."""
    name = MONTHS[month - 1].upper()
    ws.cell(row=row, column=1, value=f"{name} {year}").font = Font(
        bold=True, size=13, color=NAVY)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=HOURS_COL)
    ws.row_dimensions[row].height = 22

    _block_header(ws, row + 1)
    r = row + 2

    for week in calendar.Calendar(firstweekday=0).monthdatescalendar(year, month):
        in_month = [d for d in week if d.month == month]
        for i, d in enumerate(week, start=1):
            if d.month != month:
                for rr in (r, r + 1, r + 2):
                    ws.cell(row=rr, column=i).fill = PatternFill(
                        "solid", fgColor=OUTSIDE_FILL)
                continue
            rr_rota = rota_row(d, year)
            n = ws.cell(row=r, column=i, value=d.day)
            n.font = Font(bold=True, size=11,
                          color="B03030" if d in holidays else "000000")
            n.fill = PatternFill("solid", fgColor=_day_fill(d, holidays))
            # The asterisk marks a shift swapped with a colleague.
            code = ws.cell(row=r + 1, column=i, value=(
                f'=IF(Presenze!$D${rr_rota}="","",Presenze!$D${rr_rota}'
                f'&IF(Presenze!$F${rr_rota}="","","*"))'))
            code.font = Font(bold=True, size=12)
            times = ws.cell(row=r + 2, column=i, value=(
                f'=IF(Presenze!$D${rr_rota}="","",'
                f'IFERROR(VLOOKUP(Presenze!$D${rr_rota},{lookup},6,FALSE),""))'))
            times.font = Font(size=9, color="5A6B7D")

        areas["code"].append(f"A{r + 1}:G{r + 1}")
        areas["time"].append(f"A{r + 2}:G{r + 2}")

        # A month's days inside one week are contiguous rows in Presenze.
        first_row = rota_row(in_month[0], year)
        last_row = rota_row(in_month[-1], year)
        total = f"SUM(Presenze!$E${first_row}:$E${last_row})"
        hours = ws.cell(row=r, column=HOURS_COL, value=f'=IF({total}=0,"",{total})')
        hours.font = Font(bold=True, color=NAVY)
        ws.merge_cells(start_row=r, start_column=HOURS_COL,
                       end_row=r + 2, end_column=HOURS_COL)

        for rr, height in ((r, 18), (r + 1, 20), (r + 2, 15)):
            ws.row_dimensions[rr].height = height
            for i in range(1, HOURS_COL + 1):
                c = ws.cell(row=rr, column=i)
                c.border = BOX
                c.alignment = Alignment(horizontal="center", vertical="center")
        r += 3

    first = rota_row(date(year, month, 1), year)
    last = rota_row(date(year, month, calendar.monthrange(year, month)[1]), year)
    ws.cell(row=r, column=1, value=f"TOTALE {name}").font = Font(bold=True, color=NAVY)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    ws.cell(row=r, column=5, value="Giorni")
    ws.cell(row=r, column=6, value=f'=COUNTIF(Presenze!$E${first}:$E${last},">0")')
    ws.cell(row=r, column=7, value="Ore")
    ws.cell(row=r, column=HOURS_COL, value=f"=SUM(Presenze!$E${first}:$E${last})")
    for i in range(5, HOURS_COL + 1):
        c = ws.cell(row=r, column=i)
        c.font = Font(bold=True, color=NAVY)
        c.alignment = Alignment(horizontal="center")
        c.border = BOX
    return r + 2  # one blank row between months


def _legend(ws, year):
    """Rows 1-8: what the colours mean."""
    ws["A1"] = f"Calendario {year}"
    ws["A1"].font = Font(bold=True, size=16, color=NAVY)
    ws["A2"] = "Sola lettura: i turni si scrivono nel foglio Presenze."
    ws["A2"].font = Font(italic=True, size=10, color="5A6B7D")

    ws["A4"] = "Turni"
    ws["A4"].font = Font(bold=True, color=NAVY)
    for i, (code, description, _start, _end, _hours) in enumerate(SHIFTS, start=2):
        c = ws.cell(row=4, column=i, value=code)
        c.font = Font(bold=True, size=12)
        c.fill = PatternFill("solid", fgColor=SHIFT_FILL[code])
        c.alignment = Alignment(horizontal="center")
        c.border = BOX
        d = ws.cell(row=5, column=i, value=description)
        d.font = Font(size=9, color="5A6B7D")
        d.alignment = Alignment(horizontal="center")

    ws["A7"] = "Giorni"
    ws["A7"].font = Font(bold=True, color=NAVY)
    for i, (text, fill) in enumerate(
            [("Feriale", WEEKDAY_FILL), ("Sabato", SATURDAY_FILL),
             ("Domenica", SUNDAY_FILL), ("Festivo", HOLIDAY_FILL)], start=2):
        c = ws.cell(row=7, column=i, value=text)
        c.fill = PatternFill("solid", fgColor=fill)
        c.alignment = Alignment(horizontal="center")
        c.border = BOX

    ws["A8"] = "* accanto al codice = turno scambiato con una collega"
    ws["A8"].font = Font(italic=True, size=10, color="5A6B7D")


def build_calendar_sheet(wb, year, codes_first_row, codes_last_row):
    """Creates the Calendario sheet from the Codici sheet's data rows."""
    ws = wb.create_sheet("Calendario")
    for col in range(1, HOURS_COL):
        ws.column_dimensions[get_column_letter(col)].width = 13
    ws.column_dimensions[get_column_letter(HOURS_COL)].width = 9

    _legend(ws, year)

    holidays = italian_holidays(year)
    lookup = f"Codici!$A${codes_first_row}:$F${codes_last_row}"
    areas = {"code": [], "time": []}
    r = FIRST_BLOCK_ROW
    for month in range(1, 13):
        r = _month_block(ws, r, year, month, lookup, holidays, areas)

    # One value-based rule over every interval at once: Numbers imports these,
    # not the formula-based ones. "M" and "M*" are two different values, so
    # each code needs two rules.
    code_area = " ".join(areas["code"])
    time_area = " ".join(areas["time"])
    for entry in SHIFTS:
        code = entry[0]
        fill = PatternFill("solid", fgColor=SHIFT_FILL[code])
        for expected in (code, f"{code}*"):
            ws.conditional_formatting.add(code_area, CellIsRule(
                operator="equal", formula=[f'"{expected}"'], fill=fill))
        ws.conditional_formatting.add(time_area, CellIsRule(
            operator="equal", formula=[f'"{time_range(entry)}"'], fill=fill))
    return ws
