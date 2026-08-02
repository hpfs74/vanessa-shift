"""Foglio Calendario: dodici mini-calendari mensili, in sola lettura.

Ogni cella e' una formula che punta alla riga del foglio Presenze di quel
giorno. Qui non si scrive niente: i turni si inseriscono in Presenze.
"""

import calendar
from datetime import date

from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from comune import (BLU_SCURO, BOX, CODICI, MESI, festivi_italiani, orario_di,
                    riga_presenze)

GIORNI_BREVI = ["Lun", "Mar", "Mer", "Gio", "Ven", "Sab", "Dom"]

COL_ORE = 8              # colonna H, il totale ore della settimana
PRIMA_RIGA_BLOCCHI = 10  # sopra c'e' la legenda

# Colore della riga con il numero del giorno.
FERIALE = "FFFFFF"
SABATO = "DCE9F5"
DOMENICA = "F3DCE4"
FESTIVO = "FAE6B8"
VUOTO = "F7F9FB"     # giorni del mese precedente o successivo

# Colore delle righe codice e orario, applicato con formattazione condizionale
# perche' il codice si scrive dopo che il file e' stato generato.
COLORE_TURNO = {
    "L":  "EDF1F5",
    "M":  "FFF4CC",
    "M1": "FFE8A3",
    "P":  "FFE0C2",
    "P1": "FFCFA0",
}


def _colore_giorno(d, festivi):
    """Festivo batte domenica, domenica batte sabato."""
    if d in festivi:
        return FESTIVO
    if d.weekday() == 6:
        return DOMENICA
    if d.weekday() == 5:
        return SABATO
    return FERIALE


def _intestazione_blocco(ws, riga):
    """Riga scura Lun..Dom + Ore."""
    for i, v in enumerate(GIORNI_BREVI + ["Ore"], start=1):
        c = ws.cell(row=riga, column=i, value=v)
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=BLU_SCURO)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BOX
    ws.row_dimensions[riga].height = 20


def _blocco_mese(ws, riga, anno, mese, lookup, festivi, aree):
    """Scrive il blocco di un mese. Restituisce la prima riga libera dopo."""
    nome = MESI[mese - 1].upper()
    ws.cell(row=riga, column=1, value=f"{nome} {anno}").font = Font(
        bold=True, size=13, color=BLU_SCURO)
    ws.merge_cells(start_row=riga, start_column=1, end_row=riga, end_column=COL_ORE)
    ws.row_dimensions[riga].height = 22

    _intestazione_blocco(ws, riga + 1)
    r = riga + 2

    for settimana in calendar.Calendar(firstweekday=0).monthdatescalendar(anno, mese):
        del_mese = [d for d in settimana if d.month == mese]
        for i, d in enumerate(settimana, start=1):
            if d.month != mese:
                for rr in (r, r + 1, r + 2):
                    ws.cell(row=rr, column=i).fill = PatternFill("solid", fgColor=VUOTO)
                continue
            rp = riga_presenze(d, anno)
            n = ws.cell(row=r, column=i, value=d.day)
            n.font = Font(bold=True, size=11,
                          color="B03030" if d in festivi else "000000")
            n.fill = PatternFill("solid", fgColor=_colore_giorno(d, festivi))
            # L'asterisco segnala che quel turno e' stato scambiato con una collega.
            cod = ws.cell(row=r + 1, column=i, value=(
                f'=IF(Presenze!$D${rp}="","",Presenze!$D${rp}'
                f'&IF(Presenze!$F${rp}="","","*"))'))
            cod.font = Font(bold=True, size=12)
            ora = ws.cell(row=r + 2, column=i, value=(
                f'=IF(Presenze!$D${rp}="","",'
                f'IFERROR(VLOOKUP(Presenze!$D${rp},{lookup},6,FALSE),""))'))
            ora.font = Font(size=9, color="5A6B7D")

        aree["cod"].append(f"A{r + 1}:G{r + 1}")
        aree["ora"].append(f"A{r + 2}:G{r + 2}")

        # I giorni del mese dentro una settimana sono righe contigue in Presenze.
        r1 = riga_presenze(del_mese[0], anno)
        r2 = riga_presenze(del_mese[-1], anno)
        somma = f"SUM(Presenze!$E${r1}:$E${r2})"
        ore = ws.cell(row=r, column=COL_ORE, value=f'=IF({somma}=0,"",{somma})')
        ore.font = Font(bold=True, color=BLU_SCURO)
        ws.merge_cells(start_row=r, start_column=COL_ORE,
                       end_row=r + 2, end_column=COL_ORE)

        for rr, altezza in ((r, 18), (r + 1, 20), (r + 2, 15)):
            ws.row_dimensions[rr].height = altezza
            for i in range(1, COL_ORE + 1):
                c = ws.cell(row=rr, column=i)
                c.border = BOX
                c.alignment = Alignment(horizontal="center", vertical="center")
        r += 3

    primo = riga_presenze(date(anno, mese, 1), anno)
    ultimo = riga_presenze(date(anno, mese, calendar.monthrange(anno, mese)[1]), anno)
    ws.cell(row=r, column=1, value=f"TOTALE {nome}").font = Font(bold=True, color=BLU_SCURO)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    ws.cell(row=r, column=5, value="Giorni")
    ws.cell(row=r, column=6, value=f'=COUNTIF(Presenze!$E${primo}:$E${ultimo},">0")')
    ws.cell(row=r, column=7, value="Ore")
    ws.cell(row=r, column=COL_ORE, value=f"=SUM(Presenze!$E${primo}:$E${ultimo})")
    for i in range(5, COL_ORE + 1):
        c = ws.cell(row=r, column=i)
        c.font = Font(bold=True, color=BLU_SCURO)
        c.alignment = Alignment(horizontal="center")
        c.border = BOX
    return r + 2  # una riga vuota fra un mese e l'altro


def _legenda(ws, anno):
    """Righe 1-8: cosa vogliono dire i colori."""
    ws["A1"] = f"Calendario {anno}"
    ws["A1"].font = Font(bold=True, size=16, color=BLU_SCURO)
    ws["A2"] = "Sola lettura: i turni si scrivono nel foglio Presenze."
    ws["A2"].font = Font(italic=True, size=10, color="5A6B7D")

    ws["A4"] = "Turni"
    ws["A4"].font = Font(bold=True, color=BLU_SCURO)
    for i, (cod, desc, _ini, _fin, _ore) in enumerate(CODICI, start=2):
        c = ws.cell(row=4, column=i, value=cod)
        c.font = Font(bold=True, size=12)
        c.fill = PatternFill("solid", fgColor=COLORE_TURNO[cod])
        c.alignment = Alignment(horizontal="center")
        c.border = BOX
        d = ws.cell(row=5, column=i, value=desc)
        d.font = Font(size=9, color="5A6B7D")
        d.alignment = Alignment(horizontal="center")

    ws["A7"] = "Giorni"
    ws["A7"].font = Font(bold=True, color=BLU_SCURO)
    for i, (testo, sfondo) in enumerate(
            [("Feriale", FERIALE), ("Sabato", SABATO),
             ("Domenica", DOMENICA), ("Festivo", FESTIVO)], start=2):
        c = ws.cell(row=7, column=i, value=testo)
        c.fill = PatternFill("solid", fgColor=sfondo)
        c.alignment = Alignment(horizontal="center")
        c.border = BOX

    ws["A8"] = "* accanto al codice = turno scambiato con una collega"
    ws["A8"].font = Font(italic=True, size=10, color="5A6B7D")


def foglio_calendario(wb, anno, cod_r1, cod_r2):
    """Crea il foglio Calendario. cod_r1/cod_r2 sono le righe dati del foglio Codici."""
    ws = wb.create_sheet("Calendario")
    for col in range(1, COL_ORE):
        ws.column_dimensions[get_column_letter(col)].width = 13
    ws.column_dimensions[get_column_letter(COL_ORE)].width = 9

    _legenda(ws, anno)

    festivi = festivi_italiani(anno)
    lookup = f"Codici!$A${cod_r1}:$F${cod_r2}"
    aree = {"cod": [], "ora": []}
    r = PRIMA_RIGA_BLOCCHI
    for mese in range(1, 13):
        r = _blocco_mese(ws, r, anno, mese, lookup, festivi, aree)

    # Una regola per valore su tutti gli intervalli insieme: Numbers importa
    # queste, non quelle basate su formula. "M" e "M*" sono due valori diversi,
    # quindi servono due regole per codice.
    zona_cod = " ".join(aree["cod"])
    zona_ora = " ".join(aree["ora"])
    for voce in CODICI:
        cod = voce[0]
        sfondo = PatternFill("solid", fgColor=COLORE_TURNO[cod])
        for atteso in (cod, f"{cod}*"):
            ws.conditional_formatting.add(zona_cod, CellIsRule(
                operator="equal", formula=[f'"{atteso}"'], fill=sfondo))
        ws.conditional_formatting.add(zona_ora, CellIsRule(
            operator="equal", formula=[f'"{orario_di(voce)}"'], fill=sfondo))
    return ws
