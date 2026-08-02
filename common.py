"""Constants, styles and calendar maths shared by every sheet.

Everything written into the workbook stays in Italian: sheet names, column
headers, shift descriptions and holiday names are what Vanessa reads.
"""

from datetime import date, timedelta

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# --- Shift codes: single source of truth, the "Codici" sheet copies them ---
# (code, description, start, end, hours)
SHIFTS = [
    ("L",  "Libero",            "",      "",      0),
    ("M",  "Mattina",           "07:00", "13:00", 6),
    ("M1", "Mattina lunga",     "07:00", "14:00", 7),
    ("P",  "Pomeriggio",        "13:00", "20:00", 7),
    ("P1", "Pomeriggio lungo",  "13:00", "21:00", 8),
]
# Pie slices: only the worked shifts (Libero is 0 hours, it would not show).
WORKED_SHIFTS = [s for s in SHIFTS if s[4] > 0]

# Shift swaps. The sign says which way the favour goes:
#   +1 = the colleague owes me one, -1 = I owe her, 0 = even.
SWAP_KINDS = ["Ho coperto", "Mi ha coperto", "Scambio pari"]
SWAP_SIGN = {"Ho coperto": 1, "Mi ha coperto": -1, "Scambio pari": 0}
COLLEAGUE_ROWS = 10  # rows available in the colleague list

MONTHS = ["Gennaio", "Febbraio", "Marzo", "Aprile", "Maggio", "Giugno",
          "Luglio", "Agosto", "Settembre", "Ottobre", "Novembre", "Dicembre"]
WEEKDAYS = ["Lunedi", "Martedi", "Mercoledi", "Giovedi", "Venerdi", "Sabato", "Domenica"]

# --- Palette ---
NAVY = "1F3A5F"
GREY = "EDF1F5"        # alternating weekday rows / secondary headers
LIGHT_BLUE = "DCE9F5"  # weekend
GREEN = "E8F3EA"       # input cells
BORDER_SIDE = Side(style="thin", color="C3CDD8")
BOX = Border(left=BORDER_SIDE, right=BORDER_SIDE, top=BORDER_SIDE, bottom=BORDER_SIDE)

# First data row of the Presenze sheet (below title, note and header).
FIRST_ROTA_ROW = 5


def write_header(ws, row, values, widths=None):
    """Writes a header row in the dark style."""
    for i, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=i, value=v)
        c.font = Font(bold=True, color="FFFFFF", size=11)
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BOX
    ws.row_dimensions[row].height = 24
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w


def write_title(ws, cell, text, size=16):
    c = ws[cell]
    c.value = text
    c.font = Font(bold=True, size=size, color=NAVY)


def time_range(entry):
    """From a SHIFTS entry to the readable time range. Must produce the same
    string as the formula in Codici!F, because the colour rules compare it."""
    _code, _description, start, end, _hours = entry
    return "–" if not start else f"{start}-{end}"


def easter(year):
    """Easter Sunday, anonymous Gregorian algorithm."""
    a = year % 19
    b, c = divmod(year, 100)
    d, e = divmod(b, 4)
    f = (b + 8) // 25
    g = (b - f + 1) // 3
    h = (19 * a + b - d - g + 15) % 30
    i, k = divmod(c, 4)
    ell = (32 + 2 * e + 2 * i - h - k) % 7
    m = (a + 11 * h + 22 * ell) // 451
    month, day = divmod(h + ell - 7 * m + 114, 31)
    return date(year, month, day + 1)


def italian_holidays(year):
    """Italian national holidays: date -> name. The local patron is not here."""
    holidays = {
        date(year, 1, 1): "Capodanno",
        date(year, 1, 6): "Epifania",
        date(year, 4, 25): "Liberazione",
        date(year, 5, 1): "Festa del Lavoro",
        date(year, 6, 2): "Festa della Repubblica",
        date(year, 8, 15): "Ferragosto",
        date(year, 11, 1): "Ognissanti",
        date(year, 12, 8): "Immacolata",
        date(year, 12, 25): "Natale",
        date(year, 12, 26): "Santo Stefano",
    }
    # When Easter Monday falls on 25 April it is one day, not two:
    # the fixed holiday keeps its name.
    holidays.setdefault(easter(year) + timedelta(days=1), "Lunedi dell'Angelo")
    return holidays


def rota_row(d, year):
    """Row of the Presenze sheet holding day d."""
    return FIRST_ROTA_ROW + (d - date(year, 1, 1)).days
