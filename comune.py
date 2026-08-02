"""Costanti, stili e calcoli di calendario condivisi da tutti i fogli."""

from datetime import date, timedelta

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# --- Codici turno: unica fonte di verita', il foglio "Codici" li ricopia ---
# (codice, descrizione, inizio, fine, ore)
CODICI = [
    ("L",  "Libero",            "",      "",      0),
    ("M",  "Mattina",           "07:00", "13:00", 6),
    ("M1", "Mattina lunga",     "07:00", "14:00", 7),
    ("P",  "Pomeriggio",        "13:00", "20:00", 7),
    ("P1", "Pomeriggio lungo",  "13:00", "21:00", 8),
]
# Fette della torta: solo i turni lavorati (Libero vale 0 ore, non comparirebbe)
TURNI_LAVORATI = [c for c in CODICI if c[4] > 0]

# Scambi turno. Il segno dice da che parte sta il favore:
#   +1 = la collega mi deve un favore, -1 = lo devo io, 0 = pari e patta.
TIPI_SCAMBIO = ["Ho coperto", "Mi ha coperto", "Scambio pari"]
SEGNO_SCAMBIO = {"Ho coperto": 1, "Mi ha coperto": -1, "Scambio pari": 0}
N_COLLEGHE = 10  # righe disponibili nell'elenco colleghe

MESI = ["Gennaio", "Febbraio", "Marzo", "Aprile", "Maggio", "Giugno",
        "Luglio", "Agosto", "Settembre", "Ottobre", "Novembre", "Dicembre"]
GIORNI = ["Lunedi", "Martedi", "Mercoledi", "Giovedi", "Venerdi", "Sabato", "Domenica"]

# --- Palette ---
BLU_SCURO = "1F3A5F"
GRIGIO_INT = "EDF1F5"   # righe feriali alternate / intestazioni secondarie
AZZURRO = "DCE9F5"      # weekend
VERDE = "E8F3EA"        # celle di input
BORDO = Side(style="thin", color="C3CDD8")
BOX = Border(left=BORDO, right=BORDO, top=BORDO, bottom=BORDO)

# Prima riga di dati del foglio Presenze (sotto titolo, nota e intestazione).
PRIMA_RIGA_PRESENZE = 5


def intesta(ws, riga, valori, larghezze=None):
    """Scrive una riga di intestazione in stile scuro."""
    for i, v in enumerate(valori, start=1):
        c = ws.cell(row=riga, column=i, value=v)
        c.font = Font(bold=True, color="FFFFFF", size=11)
        c.fill = PatternFill("solid", fgColor=BLU_SCURO)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BOX
    ws.row_dimensions[riga].height = 24
    if larghezze:
        for i, w in enumerate(larghezze, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w


def titolo(ws, cella, testo, size=16):
    c = ws[cella]
    c.value = testo
    c.font = Font(bold=True, size=size, color=BLU_SCURO)


def orario_di(voce):
    """Da una voce di CODICI all'orario leggibile. Deve dare la stessa stringa
    della formula in Codici!F, perche' le regole di colore la confrontano."""
    _cod, _desc, ini, fin, _ore = voce
    return "–" if not ini else f"{ini}-{fin}"


def pasqua(anno):
    """Domenica di Pasqua, algoritmo gregoriano anonimo."""
    a = anno % 19
    b, c = divmod(anno, 100)
    d, e = divmod(b, 4)
    f = (b + 8) // 25
    g = (b - f + 1) // 3
    h = (19 * a + b - d - g + 15) % 30
    i, k = divmod(c, 4)
    ell = (32 + 2 * e + 2 * i - h - k) % 7
    m = (a + 11 * h + 22 * ell) // 451
    mese, giorno = divmod(h + ell - 7 * m + 114, 31)
    return date(anno, mese, giorno + 1)


def festivi_italiani(anno):
    """Festivi nazionali italiani: data -> nome. Il patrono locale non c'e'."""
    festivi = {
        date(anno, 1, 1): "Capodanno",
        date(anno, 1, 6): "Epifania",
        date(anno, 4, 25): "Liberazione",
        date(anno, 5, 1): "Festa del Lavoro",
        date(anno, 6, 2): "Festa della Repubblica",
        date(anno, 8, 15): "Ferragosto",
        date(anno, 11, 1): "Ognissanti",
        date(anno, 12, 8): "Immacolata",
        date(anno, 12, 25): "Natale",
        date(anno, 12, 26): "Santo Stefano",
    }
    # Se la Pasquetta cade il 25 aprile e' un giorno solo, non due: tengo il nome fisso.
    festivi.setdefault(pasqua(anno) + timedelta(days=1), "Lunedi dell'Angelo")
    return festivi


def riga_presenze(d, anno):
    """Riga del foglio Presenze che ospita il giorno d."""
    return PRIMA_RIGA_PRESENZE + (d - date(anno, 1, 1)).days
