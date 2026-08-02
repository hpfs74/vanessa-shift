#!/usr/bin/env python3
"""Genera il foglio presenze di Vanessa (.xlsx) da aprire in Numbers.

Uso:  python3 genera_presenze.py [anno] [file_output]
Es.:  python3 genera_presenze.py 2026 Presenze_Vanessa_2026.xlsx

Il codice turno si scrive nella colonna "Cod" del foglio Presenze;
le ore si calcolano da sole leggendo la tabella del foglio Codici.
"""

import calendar
import sys
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from calendario import foglio_calendario
from comune import (AZZURRO, BLU_SCURO, BOX, CODICI, GIORNI, GRIGIO_INT, MESI,
                    N_COLLEGHE, TIPI_SCAMBIO, TURNI_LAVORATI, VERDE, intesta,
                    titolo)
from stipendio import foglio_stipendio


def foglio_codici(wb):
    ws = wb.create_sheet("Codici")
    titolo(ws, "A1", "Codici turno")
    ws["A2"] = "Modifica le ore qui: tutto l'anno si ricalcola da solo."
    ws["A2"].font = Font(italic=True, size=10, color="5A6B7D")

    intesta(ws, 4, ["Cod", "Descrizione", "Inizio", "Fine", "Ore", "Orario"],
            larghezze=[10, 22, 12, 12, 10, 16])
    for i, (cod, desc, ini, fin, ore) in enumerate(CODICI):
        r = 5 + i
        for col, val in enumerate([cod, desc, ini, fin, ore], start=1):
            c = ws.cell(row=r, column=col, value=val)
            c.border = BOX
            c.alignment = Alignment(horizontal="center" if col != 2 else "left")
        # Orario leggibile, usato dal foglio Calendario: si aggiorna se cambio Inizio/Fine.
        c = ws.cell(row=r, column=6, value=f'=IF(C{r}="","–",C{r}&"-"&D{r})')
        c.border = BOX
        c.alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=1).font = Font(bold=True)
        if ore == 0:
            for col in range(1, 7):
                ws.cell(row=r, column=col).fill = PatternFill("solid", fgColor=GRIGIO_INT)
    return ws, 5, 4 + len(CODICI)  # foglio, prima riga dati, ultima riga dati


def foglio_presenze(wb, anno, cod_r1, cod_r2):
    ws = wb.create_sheet("Presenze", 0)
    titolo(ws, "A1", f"Presenze {anno}")
    ws["A2"] = "Scrivi solo il codice nella colonna Cod (L, M, M1, P, P1). Le ore si calcolano da sole."
    ws["A2"].font = Font(italic=True, size=10, color="5A6B7D")

    intesta(ws, 4, ["Data", "Giorno", "Mese", "Cod", "Ore",
                    "Cod. orig.", "Collega", "Tipo scambio", "Diff ore", "Note"],
            larghezze=[13, 13, 13, 9, 9, 11, 16, 17, 10, 30])

    d = date(anno, 1, 1)
    ultimo = date(anno, 12, 31)
    r = 5
    lookup = f"Codici!$A${cod_r1}:$E${cod_r2}"
    while d <= ultimo:
        weekend = d.weekday() >= 5
        ws.cell(row=r, column=1, value=d).number_format = "DD/MM/YYYY"
        ws.cell(row=r, column=2, value=GIORNI[d.weekday()])
        ws.cell(row=r, column=3, value=MESI[d.month - 1])
        ws.cell(row=r, column=4, value=None)
        # Ore: vuoto se non ho ancora scritto il codice, altrimenti lo cerco nei Codici
        ws.cell(row=r, column=5,
                value=f'=IF($D{r}="","",IFERROR(VLOOKUP($D{r},{lookup},5,FALSE),"?"))')
        ws.cell(row=r, column=6, value=None)   # Cod. orig. (vuoto = nessuno scambio)
        ws.cell(row=r, column=7, value=None)   # Collega
        ws.cell(row=r, column=8, value=None)   # Tipo scambio
        # Diff ore: ore fatte davvero meno ore che erano in turno.
        # Positivo = ho lavorato piu' del previsto (credito verso la collega).
        ws.cell(row=r, column=9,
                value=(f'=IF($F{r}="","",'
                       f'IFERROR(VLOOKUP($D{r},{lookup},5,FALSE),0)'
                       f'-IFERROR(VLOOKUP($F{r},{lookup},5,FALSE),0))'))
        ws.cell(row=r, column=10, value=None)  # Note

        for col in range(1, 11):
            c = ws.cell(row=r, column=col)
            c.border = BOX
            c.alignment = Alignment(horizontal="left" if col in (7, 8, 10) else "center")
            if weekend:
                c.fill = PatternFill("solid", fgColor=AZZURRO)
        # le colonne di input si distinguono a colpo d'occhio
        for col in (4, 6, 7, 8):
            ws.cell(row=r, column=col).fill = PatternFill("solid", fgColor=VERDE)
        ws.cell(row=r, column=4).font = Font(bold=True)
        r += 1
        d += timedelta(days=1)

    ultima = r - 1

    # Totale in fondo
    ws.cell(row=r + 1, column=4, value="TOTALE").font = Font(bold=True, color=BLU_SCURO)
    for col, formula in ((5, f"=SUM($E$5:$E${ultima})"), (9, f"=SUM($I$5:$I${ultima})")):
        tot = ws.cell(row=r + 1, column=col, value=formula)
        tot.font = Font(bold=True, color=BLU_SCURO)
        tot.border = BOX

    dv = DataValidation(type="list", formula1='"{}"'.format(",".join(c[0] for c in CODICI)),
                        allow_blank=True, showDropDown=False)
    ws.add_data_validation(dv)
    dv.add(f"D5:D{ultima}")
    dv.add(f"F5:F{ultima}")

    dv_tipo = DataValidation(type="list", formula1='"{}"'.format(",".join(TIPI_SCAMBIO)),
                             allow_blank=True, showDropDown=False)
    ws.add_data_validation(dv_tipo)
    dv_tipo.add(f"H5:H{ultima}")

    ws.freeze_panes = "A5"
    return ws, 5, ultima


def foglio_riepilogo(wb, anno, p_r1, p_r2):
    ws = wb.create_sheet("Riepilogo")
    titolo(ws, "A1", f"Riepilogo {anno}")

    codici = [c[0] for c in CODICI]
    mese_rng = f"Presenze!$C${p_r1}:$C${p_r2}"
    cod_rng = f"Presenze!$D${p_r1}:$D${p_r2}"
    ore_rng = f"Presenze!$E${p_r1}:$E${p_r2}"

    # --- Tabella mensile ---
    testata = ["Mese"] + codici + ["Giorni lavorati", "Ore totali"]
    # colonna A larga: ospita anche le etichette lunghe delle tabelle dei grafici
    intesta(ws, 3, testata, larghezze=[30] + [8] * len(codici) + [17, 13])
    for i, mese in enumerate(MESI):
        r = 4 + i
        ws.cell(row=r, column=1, value=mese).font = Font(bold=True)
        for j, cod in enumerate(codici):
            col = 2 + j
            ws.cell(row=r, column=col,
                    value=f'=COUNTIFS({mese_rng},$A{r},{cod_rng},"{cod}")')
        col_gg = 2 + len(codici)
        ws.cell(row=r, column=col_gg,
                value=f'=COUNTIFS({mese_rng},$A{r},{ore_rng},">0")')
        ws.cell(row=r, column=col_gg + 1,
                value=f"=SUMIFS({ore_rng},{mese_rng},$A{r})")
        for col in range(1, col_gg + 2):
            c = ws.cell(row=r, column=col)
            c.border = BOX
            if col > 1:
                c.alignment = Alignment(horizontal="center")
            if i % 2 == 0:
                c.fill = PatternFill("solid", fgColor=GRIGIO_INT)

    r_tot = 4 + len(MESI)
    ws.cell(row=r_tot, column=1, value="TOTALE ANNO")
    for col in range(2, 2 + len(codici) + 2):
        L = get_column_letter(col)
        ws.cell(row=r_tot, column=col, value=f"=SUM({L}4:{L}{r_tot - 1})")
    for col in range(1, 2 + len(codici) + 2):
        c = ws.cell(row=r_tot, column=col)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=BLU_SCURO)
        c.border = BOX
        if col > 1:
            c.alignment = Alignment(horizontal="center")

    # --- Dati torta 1: ore per turno (i giorni liberi valgono 0 ore, non entrano) ---
    r0 = r_tot + 3
    titolo(ws, f"A{r0 - 1}", "Ore per turno", size=12)
    intesta(ws, r0, ["Turno", "Ore"])
    for i, (cod, desc, ini, fin, _ore) in enumerate(TURNI_LAVORATI):
        r = r0 + 1 + i
        ws.cell(row=r, column=1, value=f"{desc} ({ini}-{fin})").border = BOX
        c = ws.cell(row=r, column=2, value=f'=SUMIFS({ore_rng},{cod_rng},"{cod}")')
        c.border = BOX
        c.alignment = Alignment(horizontal="center")
    ore_first, ore_last = r0 + 1, r0 + len(TURNI_LAVORATI)

    # --- Dati torta 2: giorni per codice (qui il Libero si vede) ---
    r1 = ore_last + 3
    titolo(ws, f"A{r1 - 1}", "Giorni per codice", size=12)
    intesta(ws, r1, ["Codice", "Giorni"])
    for i, (cod, desc, _i, _f, _o) in enumerate(CODICI):
        r = r1 + 1 + i
        ws.cell(row=r, column=1, value=f"{cod} - {desc}").border = BOX
        c = ws.cell(row=r, column=2, value=f'=COUNTIF({cod_rng},"{cod}")')
        c.border = BOX
        c.alignment = Alignment(horizontal="center")
    gg_first, gg_last = r1 + 1, r1 + len(CODICI)

    # --- Grafici ---
    def torta(titolo_g, first, last, ancora):
        ch = PieChart()
        ch.title = titolo_g
        ch.height, ch.width = 9.5, 15
        ch.add_data(Reference(ws, min_col=2, min_row=first - 1, max_row=last), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=first, max_row=last))
        ch.dataLabels = DataLabelList()
        ch.dataLabels.showVal = True
        ch.dataLabels.showPercent = True
        ws.add_chart(ch, ancora)
        return ch

    # ancorati nella colonna J, ben distanziati: altrimenti si sovrappongono fra loro
    torta("Ore per turno", ore_first, ore_last, "J3")
    torta("Giorni per codice", gg_first, gg_last, "J25")
    return ws


def foglio_scambi(wb, anno, p_r1, p_r2):
    """Saldo dei favori: chi mi deve un turno e chi devo io."""
    ws = wb.create_sheet("Scambi")
    titolo(ws, "A1", f"Scambi turno {anno}")
    ws["A2"] = ("Scrivi qui i nomi delle colleghe, poi usali nella colonna Collega del foglio Presenze. "
                "Saldo positivo = lei deve un favore a te.")
    ws["A2"].font = Font(italic=True, size=10, color="5A6B7D")

    coll_rng = f"Presenze!$G${p_r1}:$G${p_r2}"
    tipo_rng = f"Presenze!$H${p_r1}:$H${p_r2}"
    diff_rng = f"Presenze!$I${p_r1}:$I${p_r2}"

    intesta(ws, 4, ["Collega"] + TIPI_SCAMBIO + ["Saldo favori", "Saldo ore"],
            larghezze=[30] + [14] * len(TIPI_SCAMBIO) + [14, 13])

    r0 = 5
    for i in range(N_COLLEGHE):
        r = r0 + i
        nome = ws.cell(row=r, column=1, value=None)
        nome.fill = PatternFill("solid", fgColor=VERDE)
        nome.font = Font(bold=True)
        for j, tipo in enumerate(TIPI_SCAMBIO):
            ws.cell(row=r, column=2 + j,
                    value=f'=IF($A{r}="","",COUNTIFS({coll_rng},$A{r},{tipo_rng},"{tipo}"))')
        # saldo favori: quante volte ho coperto io meno quante volte mi hanno coperto
        c_piu = get_column_letter(2 + TIPI_SCAMBIO.index("Ho coperto"))
        c_meno = get_column_letter(2 + TIPI_SCAMBIO.index("Mi ha coperto"))
        ws.cell(row=r, column=2 + len(TIPI_SCAMBIO),
                value=f'=IF($A{r}="","",{c_piu}{r}-{c_meno}{r})')
        ws.cell(row=r, column=3 + len(TIPI_SCAMBIO),
                value=f'=IF($A{r}="","",SUMIFS({diff_rng},{coll_rng},$A{r}))')
        for col in range(1, 4 + len(TIPI_SCAMBIO)):
            c = ws.cell(row=r, column=col)
            c.border = BOX
            if col > 1:
                c.alignment = Alignment(horizontal="center")
    r_last = r0 + N_COLLEGHE - 1

    # Rete di sicurezza: se un nome nelle Presenze non e' in elenco (refuso), qui salta fuori.
    r_chk = r_last + 2
    ws.cell(row=r_chk, column=1, value="Scambi registrati").font = Font(bold=True)
    ws.cell(row=r_chk, column=2, value=f'=COUNTIF({tipo_rng},"<>")')
    ws.cell(row=r_chk + 1, column=1, value="di cui abbinati a un nome").font = Font(bold=True)
    somma = "+".join(f"SUM({get_column_letter(2 + j)}{r0}:{get_column_letter(2 + j)}{r_last})"
                     for j in range(len(TIPI_SCAMBIO)))
    ws.cell(row=r_chk + 1, column=2, value=f"={somma}")
    ws.cell(row=r_chk + 2, column=1, value="Da controllare (refusi)").font = Font(bold=True, color="B03030")
    c = ws.cell(row=r_chk + 2, column=2, value=f"=B{r_chk}-B{r_chk + 1}")
    c.font = Font(bold=True, color="B03030")
    for rr in (r_chk, r_chk + 1, r_chk + 2):
        ws.cell(row=rr, column=2).border = BOX
        ws.cell(row=rr, column=2).alignment = Alignment(horizontal="center")

    # Grafico: saldo ore per collega (sopra lo zero = credito)
    ch = BarChart()
    ch.type = "col"
    ch.title = "Saldo ore per collega"
    ch.height, ch.width = 9, 16
    ch.legend = None
    ch.y_axis.title = "Ore"
    ch.add_data(Reference(ws, min_col=3 + len(TIPI_SCAMBIO), min_row=4, max_row=r_last),
                titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=r0, max_row=r_last))
    ws.add_chart(ch, "I4")
    return ws


def main():
    anno = int(sys.argv[1]) if len(sys.argv) > 1 else 2026
    out = sys.argv[2] if len(sys.argv) > 2 else f"Presenze_Vanessa_{anno}.xlsx"

    wb = Workbook()
    wb.remove(wb.active)
    _, cod_r1, cod_r2 = foglio_codici(wb)
    _, p_r1, p_r2 = foglio_presenze(wb, anno, cod_r1, cod_r2)
    foglio_calendario(wb, anno, cod_r1, cod_r2)
    foglio_riepilogo(wb, anno, p_r1, p_r2)
    foglio_scambi(wb, anno, p_r1, p_r2)
    foglio_stipendio(wb, anno, p_r1, p_r2)
    wb.move_sheet("Presenze", offset=-wb.sheetnames.index("Presenze"))
    wb.save(out)
    giorni = 366 if calendar.isleap(anno) else 365
    print(f"Creato {out} - {giorni} giorni, fogli: {', '.join(wb.sheetnames)}")


if __name__ == "__main__":
    main()
