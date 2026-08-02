#!/usr/bin/env python3
"""Generates Vanessa's rota workbook (.xlsx) to be opened in Numbers.

Usage:  python3 generate_rota.py [year] [output_file]
e.g.:   python3 generate_rota.py 2026 Presenze_Vanessa_2026.xlsx

The shift code is typed into the "Cod" column of the Presenze sheet;
the hours work themselves out by reading the Codici sheet's table.

Sheet names, headings and every string written into the workbook stay in
Italian: they are what Vanessa reads.
"""

import calendar
import sys
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from calendar_sheet import build_calendar_sheet
from common import (BOX, COLLEAGUE_ROWS, GREEN, GREY, LIGHT_BLUE, MONTHS, NAVY,
                    SHIFTS, SWAP_KINDS, WEEKDAYS, WORKED_SHIFTS, write_header,
                    write_title)
from pay_sheet import build_pay_sheet


def build_codes_sheet(wb):
    ws = wb.create_sheet("Codici")
    write_title(ws, "A1", "Codici turno")
    ws["A2"] = "Modifica le ore qui: tutto l'anno si ricalcola da solo."
    ws["A2"].font = Font(italic=True, size=10, color="5A6B7D")

    write_header(ws, 4, ["Cod", "Descrizione", "Inizio", "Fine", "Ore", "Orario"],
                 widths=[10, 22, 12, 12, 10, 16])
    for i, (code, description, start, end, hours) in enumerate(SHIFTS):
        r = 5 + i
        for col, value in enumerate([code, description, start, end, hours], start=1):
            c = ws.cell(row=r, column=col, value=value)
            c.border = BOX
            c.alignment = Alignment(horizontal="center" if col != 2 else "left")
        # Readable time range, used by the Calendario sheet: it follows any
        # change to Inizio/Fine.
        c = ws.cell(row=r, column=6, value=f'=IF(C{r}="","–",C{r}&"-"&D{r})')
        c.border = BOX
        c.alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=1).font = Font(bold=True)
        if hours == 0:
            for col in range(1, 7):
                ws.cell(row=r, column=col).fill = PatternFill("solid", fgColor=GREY)
    return ws, 5, 4 + len(SHIFTS)  # sheet, first data row, last data row


def build_rota_sheet(wb, year, codes_first_row, codes_last_row):
    ws = wb.create_sheet("Presenze", 0)
    write_title(ws, "A1", f"Presenze {year}")
    ws["A2"] = ("Scrivi solo il codice nella colonna Cod (L, M, M1, P, P1). "
                "Le ore si calcolano da sole.")
    ws["A2"].font = Font(italic=True, size=10, color="5A6B7D")

    write_header(ws, 4, ["Data", "Giorno", "Mese", "Cod", "Ore",
                         "Cod. orig.", "Collega", "Tipo scambio", "Diff ore", "Note"],
                 widths=[13, 13, 13, 9, 9, 11, 16, 17, 10, 30])

    d = date(year, 1, 1)
    last_day = date(year, 12, 31)
    r = 5
    lookup = f"Codici!$A${codes_first_row}:$E${codes_last_row}"
    while d <= last_day:
        is_weekend = d.weekday() >= 5
        ws.cell(row=r, column=1, value=d).number_format = "DD/MM/YYYY"
        ws.cell(row=r, column=2, value=WEEKDAYS[d.weekday()])
        ws.cell(row=r, column=3, value=MONTHS[d.month - 1])
        ws.cell(row=r, column=4, value=None)
        # Hours: blank until the code is typed, otherwise looked up in Codici.
        ws.cell(row=r, column=5,
                value=f'=IF($D{r}="","",IFERROR(VLOOKUP($D{r},{lookup},5,FALSE),"?"))')
        ws.cell(row=r, column=6, value=None)   # Cod. orig. (blank = no swap)
        ws.cell(row=r, column=7, value=None)   # Collega
        ws.cell(row=r, column=8, value=None)   # Tipo scambio
        # Diff ore: hours actually worked minus hours originally rostered.
        # Positive = worked more than planned (credit towards the colleague).
        ws.cell(row=r, column=9,
                value=(f'=IF($F{r}="","",'
                       f'IFERROR(VLOOKUP($D{r},{lookup},5,FALSE),0)'
                       f'-IFERROR(VLOOKUP($F{r},{lookup},5,FALSE),0))'))
        ws.cell(row=r, column=10, value=None)  # Note

        for col in range(1, 11):
            c = ws.cell(row=r, column=col)
            c.border = BOX
            c.alignment = Alignment(horizontal="left" if col in (7, 8, 10) else "center")
            if is_weekend:
                c.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        # The input columns stand out at a glance.
        for col in (4, 6, 7, 8):
            ws.cell(row=r, column=col).fill = PatternFill("solid", fgColor=GREEN)
        ws.cell(row=r, column=4).font = Font(bold=True)
        r += 1
        d += timedelta(days=1)

    last_row = r - 1

    # Total at the bottom
    ws.cell(row=r + 1, column=4, value="TOTALE").font = Font(bold=True, color=NAVY)
    for col, formula in ((5, f"=SUM($E$5:$E${last_row})"),
                         (9, f"=SUM($I$5:$I${last_row})")):
        total = ws.cell(row=r + 1, column=col, value=formula)
        total.font = Font(bold=True, color=NAVY)
        total.border = BOX

    codes_dv = DataValidation(
        type="list", formula1='"{}"'.format(",".join(s[0] for s in SHIFTS)),
        allow_blank=True, showDropDown=False)
    ws.add_data_validation(codes_dv)
    codes_dv.add(f"D5:D{last_row}")
    codes_dv.add(f"F5:F{last_row}")

    swap_dv = DataValidation(type="list", formula1='"{}"'.format(",".join(SWAP_KINDS)),
                             allow_blank=True, showDropDown=False)
    ws.add_data_validation(swap_dv)
    swap_dv.add(f"H5:H{last_row}")

    ws.freeze_panes = "A5"
    return ws, 5, last_row


def build_summary_sheet(wb, year, rota_first_row, rota_last_row):
    ws = wb.create_sheet("Riepilogo")
    write_title(ws, "A1", f"Riepilogo {year}")

    codes = [s[0] for s in SHIFTS]
    month_range = f"Presenze!$C${rota_first_row}:$C${rota_last_row}"
    code_range = f"Presenze!$D${rota_first_row}:$D${rota_last_row}"
    hours_range = f"Presenze!$E${rota_first_row}:$E${rota_last_row}"

    # --- Monthly table ---
    heading = ["Mese"] + codes + ["Giorni lavorati", "Ore totali"]
    # Column A is wide: it also holds the long labels of the chart tables.
    write_header(ws, 3, heading, widths=[30] + [8] * len(codes) + [17, 13])
    for i, month in enumerate(MONTHS):
        r = 4 + i
        ws.cell(row=r, column=1, value=month).font = Font(bold=True)
        for j, code in enumerate(codes):
            col = 2 + j
            ws.cell(row=r, column=col,
                    value=f'=COUNTIFS({month_range},$A{r},{code_range},"{code}")')
        days_col = 2 + len(codes)
        ws.cell(row=r, column=days_col,
                value=f'=COUNTIFS({month_range},$A{r},{hours_range},">0")')
        ws.cell(row=r, column=days_col + 1,
                value=f"=SUMIFS({hours_range},{month_range},$A{r})")
        for col in range(1, days_col + 2):
            c = ws.cell(row=r, column=col)
            c.border = BOX
            if col > 1:
                c.alignment = Alignment(horizontal="center")
            if i % 2 == 0:
                c.fill = PatternFill("solid", fgColor=GREY)

    total_row = 4 + len(MONTHS)
    ws.cell(row=total_row, column=1, value="TOTALE ANNO")
    for col in range(2, 2 + len(codes) + 2):
        letter = get_column_letter(col)
        ws.cell(row=total_row, column=col,
                value=f"=SUM({letter}4:{letter}{total_row - 1})")
    for col in range(1, 2 + len(codes) + 2):
        c = ws.cell(row=total_row, column=col)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.border = BOX
        if col > 1:
            c.alignment = Alignment(horizontal="center")

    # --- Pie data 1: hours per shift (days off are 0 hours, they stay out) ---
    r0 = total_row + 3
    write_title(ws, f"A{r0 - 1}", "Ore per turno", size=12)
    write_header(ws, r0, ["Turno", "Ore"])
    for i, (code, description, start, end, _hours) in enumerate(WORKED_SHIFTS):
        r = r0 + 1 + i
        ws.cell(row=r, column=1, value=f"{description} ({start}-{end})").border = BOX
        c = ws.cell(row=r, column=2, value=f'=SUMIFS({hours_range},{code_range},"{code}")')
        c.border = BOX
        c.alignment = Alignment(horizontal="center")
    hours_first, hours_last = r0 + 1, r0 + len(WORKED_SHIFTS)

    # --- Pie data 2: days per code (here Libero does show) ---
    r1 = hours_last + 3
    write_title(ws, f"A{r1 - 1}", "Giorni per codice", size=12)
    write_header(ws, r1, ["Codice", "Giorni"])
    for i, (code, description, _s, _e, _h) in enumerate(SHIFTS):
        r = r1 + 1 + i
        ws.cell(row=r, column=1, value=f"{code} - {description}").border = BOX
        c = ws.cell(row=r, column=2, value=f'=COUNTIF({code_range},"{code}")')
        c.border = BOX
        c.alignment = Alignment(horizontal="center")
    days_first, days_last = r1 + 1, r1 + len(SHIFTS)

    # --- Charts ---
    def pie(title, first, last, anchor):
        ch = PieChart()
        ch.title = title
        ch.height, ch.width = 9.5, 15
        ch.add_data(Reference(ws, min_col=2, min_row=first - 1, max_row=last),
                    titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=first, max_row=last))
        ch.dataLabels = DataLabelList()
        ch.dataLabels.showVal = True
        ch.dataLabels.showPercent = True
        ws.add_chart(ch, anchor)
        return ch

    # Anchored in column J and well spaced: otherwise they overlap each other.
    pie("Ore per turno", hours_first, hours_last, "J3")
    pie("Giorni per codice", days_first, days_last, "J25")
    return ws


def build_swaps_sheet(wb, year, rota_first_row, rota_last_row):
    """Balance of favours: who owes me a shift and who I owe."""
    ws = wb.create_sheet("Scambi")
    write_title(ws, "A1", f"Scambi turno {year}")
    ws["A2"] = ("Scrivi qui i nomi delle colleghe, poi usali nella colonna Collega "
                "del foglio Presenze. Saldo positivo = lei deve un favore a te.")
    ws["A2"].font = Font(italic=True, size=10, color="5A6B7D")

    colleague_range = f"Presenze!$G${rota_first_row}:$G${rota_last_row}"
    kind_range = f"Presenze!$H${rota_first_row}:$H${rota_last_row}"
    diff_range = f"Presenze!$I${rota_first_row}:$I${rota_last_row}"

    write_header(ws, 4, ["Collega"] + SWAP_KINDS + ["Saldo favori", "Saldo ore"],
                 widths=[30] + [14] * len(SWAP_KINDS) + [14, 13])

    r0 = 5
    for i in range(COLLEAGUE_ROWS):
        r = r0 + i
        name = ws.cell(row=r, column=1, value=None)
        name.fill = PatternFill("solid", fgColor=GREEN)
        name.font = Font(bold=True)
        for j, kind in enumerate(SWAP_KINDS):
            ws.cell(row=r, column=2 + j,
                    value=(f'=IF($A{r}="","",'
                           f'COUNTIFS({colleague_range},$A{r},{kind_range},"{kind}"))'))
        # Favour balance: how often I covered minus how often I was covered.
        covered_col = get_column_letter(2 + SWAP_KINDS.index("Ho coperto"))
        owed_col = get_column_letter(2 + SWAP_KINDS.index("Mi ha coperto"))
        ws.cell(row=r, column=2 + len(SWAP_KINDS),
                value=f'=IF($A{r}="","",{covered_col}{r}-{owed_col}{r})')
        ws.cell(row=r, column=3 + len(SWAP_KINDS),
                value=f'=IF($A{r}="","",SUMIFS({diff_range},{colleague_range},$A{r}))')
        for col in range(1, 4 + len(SWAP_KINDS)):
            c = ws.cell(row=r, column=col)
            c.border = BOX
            if col > 1:
                c.alignment = Alignment(horizontal="center")
    last_row = r0 + COLLEAGUE_ROWS - 1

    # Safety net: a name in Presenze that is not in the list (a typo) shows up here.
    check_row = last_row + 2
    ws.cell(row=check_row, column=1, value="Scambi registrati").font = Font(bold=True)
    ws.cell(row=check_row, column=2, value=f'=COUNTIF({kind_range},"<>")')
    ws.cell(row=check_row + 1, column=1,
            value="di cui abbinati a un nome").font = Font(bold=True)
    matched = "+".join(
        f"SUM({get_column_letter(2 + j)}{r0}:{get_column_letter(2 + j)}{last_row})"
        for j in range(len(SWAP_KINDS)))
    ws.cell(row=check_row + 1, column=2, value=f"={matched}")
    ws.cell(row=check_row + 2, column=1,
            value="Da controllare (refusi)").font = Font(bold=True, color="B03030")
    c = ws.cell(row=check_row + 2, column=2, value=f"=B{check_row}-B{check_row + 1}")
    c.font = Font(bold=True, color="B03030")
    for rr in (check_row, check_row + 1, check_row + 2):
        ws.cell(row=rr, column=2).border = BOX
        ws.cell(row=rr, column=2).alignment = Alignment(horizontal="center")

    # Chart: hour balance per colleague (above zero = credit).
    chart = BarChart()
    chart.type = "col"
    chart.title = "Saldo ore per collega"
    chart.height, chart.width = 9, 16
    chart.legend = None
    chart.y_axis.title = "Ore"
    chart.add_data(
        Reference(ws, min_col=3 + len(SWAP_KINDS), min_row=4, max_row=last_row),
        titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=r0, max_row=last_row))
    ws.add_chart(chart, "I4")
    return ws


def main():
    year = int(sys.argv[1]) if len(sys.argv) > 1 else 2026
    out = sys.argv[2] if len(sys.argv) > 2 else f"Presenze_Vanessa_{year}.xlsx"

    wb = Workbook()
    wb.remove(wb.active)
    _, codes_first, codes_last = build_codes_sheet(wb)
    _, rota_first, rota_last = build_rota_sheet(wb, year, codes_first, codes_last)
    build_calendar_sheet(wb, year, codes_first, codes_last)
    build_summary_sheet(wb, year, rota_first, rota_last)
    build_swaps_sheet(wb, year, rota_first, rota_last)
    build_pay_sheet(wb, year, rota_first, rota_last)
    wb.move_sheet("Presenze", offset=-wb.sheetnames.index("Presenze"))
    wb.save(out)
    days = 366 if calendar.isleap(year) else 365
    print(f"Creato {out} - {days} giorni, fogli: {', '.join(wb.sheetnames)}")


if __name__ == "__main__":
    main()
