"""Stipendio sheet: monthly pay simulation.

Hours come from the Presenze sheet. The rate and the premiums are typed by
the user into the green cells: while they are empty the table stays empty.

Sheet name, labels and notes stay in Italian: Vanessa reads them.
"""

from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from common import (BOX, GREEN, GREY, MONTHS, NAVY, italian_holidays,
                    rota_row, write_header, write_title)

SETTINGS_ROW = 5
TABLE_ROW = 15                    # header of the monthly table
FIRST_MONTH_ROW = 16
TOTAL_ROW = FIRST_MONTH_ROW + 12  # 28
CHECK_ROW = TOTAL_ROW + 2         # 30

COLUMNS = ["Mese", "Ore ord.", "Ore sab", "Ore dom", "Ore fest",
           "Lordo base", "Magg. sab", "Magg. dom", "Magg. fest",
           "Rateo 13a", "Lordo totale", "Netto stimato"]

SETTINGS = [
    ("Tariffa oraria lorda", "€ #,##0.00"),
    ("Maggiorazione sabato", "0.00%"),
    ("Maggiorazione domenica", "0.00%"),
    ("Maggiorazione festivo", "0.00%"),
    ("Rateo 13a", "0.00%"),
    ("Coefficiente netto/lordo", "0.00%"),
]


def _settings_block(ws, year):
    """Rows 1-13: the cells to fill in by hand, and the caveats."""
    write_title(ws, "A1", f"Simulazione stipendio {year}")
    ws["A2"] = ("Compila le celle verdi: la tabella qui sotto si calcola da sola. "
                "CCNL Cooperative Sociali, OSS livello C1.")
    ws["A2"].font = Font(italic=True, size=10, color="5A6B7D")

    write_header(ws, 4, ["Parametro", "Valore"])
    for i, (label, number_format) in enumerate(SETTINGS):
        r = 5 + i
        e = ws.cell(row=r, column=1, value=label)
        e.font = Font(bold=True)
        e.border = BOX
        c = ws.cell(row=r, column=2)
        # The accrual is the only one with a sensible default: 1/12 of the gross.
        c.value = "=1/12" if label == "Rateo 13a" else None
        c.number_format = number_format
        c.fill = PatternFill("solid", fgColor=GREEN)
        c.alignment = Alignment(horizontal="center")
        c.border = BOX

    ws["A12"] = ("La tariffa oraria si legge sul contratto o sulla busta paga. "
                 "Il coefficiente netto/lordo si ottiene dividendo il netto di una "
                 "busta paga vera per il suo lordo.")
    ws["A12"].font = Font(italic=True, size=10, color="5A6B7D")
    ws["A13"] = ("Non calcola: straordinari, notturno, scatti di anzianita', TFR, "
                 "conguagli, addizionali regionali e comunali.")
    ws["A13"].font = Font(italic=True, size=10, color="5A6B7D")


def _sum_of_rows(rows):
    """Explicit sum of Presenze cells, or 0 when there are none.

    N() reads a still-empty cell as zero hours: here, unlike the euro amounts,
    emptiness is a true fact (no shift entered), not missing information, so
    zero is the correct reading rather than a misleading one. Without N(),
    adding two empty-text cells with "+" yields #VALUE! instead of 0.
    """
    if not rows:
        return 0
    return "=" + "+".join(f"N(Presenze!$E${r})" for r in rows)


def build_pay_sheet(wb, year, rota_first_row, rota_last_row):
    ws = wb.create_sheet("Stipendio")
    _settings_block(ws, year)
    holidays = italian_holidays(year)
    hours = f"Presenze!$E${rota_first_row}:$E${rota_last_row}"
    months = f"Presenze!$C${rota_first_row}:$C${rota_last_row}"
    weekdays = f"Presenze!$B${rota_first_row}:$B${rota_last_row}"

    write_header(ws, TABLE_ROW, COLUMNS, widths=[14] + [10] * 4 + [13] * 7)

    for i, name in enumerate(MONTHS):
        r = FIRST_MONTH_ROW + i
        month = i + 1
        in_month = sorted(d for d in holidays if d.month == month)
        holiday_rows = [rota_row(d, year) for d in in_month]
        saturday_holidays = [rota_row(d, year) for d in in_month if d.weekday() == 5]
        sunday_holidays = [rota_row(d, year) for d in in_month if d.weekday() == 6]

        ws.cell(row=r, column=1, value=name).font = Font(bold=True)
        # Precedence: holiday, then Sunday, then Saturday. One bucket per hour.
        ws.cell(row=r, column=5, value=_sum_of_rows(holiday_rows))
        # N() on the subtractions for the same reason as _sum_of_rows: a
        # Saturday or Sunday holiday not yet worked must not break the
        # calculation with #VALUE!, it must simply subtract zero.
        ws.cell(row=r, column=3, value=(
            f'=SUMIFS({hours},{months},$A{r},{weekdays},"Sabato")'
            + "".join(f"-N(Presenze!$E${n})" for n in saturday_holidays)))
        ws.cell(row=r, column=4, value=(
            f'=SUMIFS({hours},{months},$A{r},{weekdays},"Domenica")'
            + "".join(f"-N(Presenze!$E${n})" for n in sunday_holidays)))
        ws.cell(row=r, column=2, value=(
            f'=SUMIFS({hours},{months},$A{r})-C{r}-D{r}-E{r}'))

        # Every hour at the base rate, then premiums only on sat/sun/holiday.
        # Each premium stays blank until its own percentage is set, not just
        # the rate: otherwise "blank times number" gives 0, another figure
        # that looks like a result and is only a missing parameter.
        ws.cell(row=r, column=6, value=f'=IF($B$5="","",(B{r}+C{r}+D{r}+E{r})*$B$5)')
        ws.cell(row=r, column=7, value=f'=IF(OR($B$5="",$B$6=""),"",C{r}*$B$5*$B$6)')
        ws.cell(row=r, column=8, value=f'=IF(OR($B$5="",$B$7=""),"",D{r}*$B$5*$B$7)')
        ws.cell(row=r, column=9, value=f'=IF(OR($B$5="",$B$8=""),"",E{r}*$B$5*$B$8)')
        # The accrual and the gross total add columns that may be blank (empty
        # text), not zero: "blank"+number would give #VALUE!, so each checks
        # that its addends upstream are already resolved before summing them.
        # G/H/I being blank already means "B5, B6, B7 or B8 is missing":
        # checking them is enough, no need to repeat the parameters.
        ws.cell(row=r, column=10, value=(
            f'=IF(OR(G{r}="",H{r}="",I{r}="",$B$9=""),"",(F{r}+G{r}+H{r}+I{r})*$B$9)'))
        ws.cell(row=r, column=11, value=f'=IF(J{r}="","",F{r}+G{r}+H{r}+I{r}+J{r})')
        ws.cell(row=r, column=12, value=f'=IF(OR(K{r}="",$B$10=""),"",K{r}*$B$10)')
        for col in range(6, 13):
            ws.cell(row=r, column=col).number_format = '€ #,##0.00'

        for col in range(1, len(COLUMNS) + 1):
            c = ws.cell(row=r, column=col)
            c.border = BOX
            if col > 1:
                c.alignment = Alignment(horizontal="center")
            if i % 2 == 0:
                c.fill = PatternFill("solid", fgColor=GREY)

    # Yearly total
    ws.cell(row=TOTAL_ROW, column=1, value="TOTALE ANNO")
    for col in range(2, len(COLUMNS) + 1):
        letter = get_column_letter(col)
        span = f"{letter}{FIRST_MONTH_ROW}:{letter}{TOTAL_ROW - 1}"
        if col >= 6:
            # SUM ignores empty-text cells and would give 0 on a year that is
            # not configured yet: a zero-euro total that looks like a real
            # result. January acts as the sentinel: if it is blank so are the
            # other eleven, same switch.
            value = f'=IF({letter}{FIRST_MONTH_ROW}="","",SUM({span}))'
        else:
            value = f"=SUM({span})"
        ws.cell(row=TOTAL_ROW, column=col, value=value)
    for col in range(1, len(COLUMNS) + 1):
        c = ws.cell(row=TOTAL_ROW, column=col)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.border = BOX
        if col > 1:
            c.alignment = Alignment(horizontal="center")
        if col >= 6:
            c.number_format = '€ #,##0.00'

    # Safety net: B is built by difference (SUMIFS(month)-C-D-E), so B+C+D+E is
    # identically equal to SUM(hours): a double count across
    # saturday/sunday/holiday would never move it off zero, which makes it no
    # check at all. A day counted twice does drain that month's ordinary hours,
    # which can fall below zero: that is where a bucketing defect really shows.
    ws.cell(row=CHECK_ROW, column=1,
            value="Controllo ore ordinarie (minimo mensile, deve essere >= 0)").font = Font(
        bold=True, color="B03030")
    c = ws.cell(row=CHECK_ROW, column=2,
                value=f"=MIN(B{FIRST_MONTH_ROW}:B{TOTAL_ROW - 1})")
    c.font = Font(bold=True, color="B03030")
    c.border = BOX
    c.alignment = Alignment(horizontal="center")
    ws.conditional_formatting.add(
        f"B{CHECK_ROW}",
        CellIsRule(operator="lessThan", formula=["0"],
                   font=Font(bold=True, color="FFFFFF"),
                   fill=PatternFill("solid", fgColor="B03030")))

    chart = BarChart()
    chart.type = "col"
    chart.title = "Lordo e netto stimato per mese"
    chart.height, chart.width = 9, 18
    chart.y_axis.title = "Euro"
    chart.add_data(Reference(ws, min_col=11, max_col=12,
                             min_row=TABLE_ROW, max_row=TOTAL_ROW - 1),
                   titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1,
                                   min_row=FIRST_MONTH_ROW, max_row=TOTAL_ROW - 1))
    ws.add_chart(chart, "N4")
    return ws
