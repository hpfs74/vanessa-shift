"""Tests for the Stipendio sheet."""

import re
import shutil
import subprocess
import sys
from datetime import date, timedelta

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

import common
import pay_sheet
from conftest import YEAR, ROOT

MONTH_ROWS = {m: 16 + i for i, m in enumerate(common.MONTHS)}

# --- Patterns for the only formula shapes pay_sheet.py can produce in
# columns B/C/D/E. The partition test below needs them: given a year, they
# match any Presenze!$E$p_r1:$E$p_r2 range rather than a fixed 5:369, so the
# test stays valid when the year changes. ---
_RE_HOLIDAY_TERM = re.compile(r'^N\(Presenze!\$E\$(\d+)\)$')
_RE_SUBTRACTION = re.compile(r'-N\(Presenze!\$E\$(\d+)\)')


def _year_bounds(year):
    """First and last Presenze rows for the year, plus the (row, date) list."""
    p_r1 = common.rota_row(date(year, 1, 1), year)
    p_r2 = common.rota_row(date(year, 12, 31), year)
    days, d, r = [], date(year, 1, 1), p_r1
    while r <= p_r2:
        days.append((r, d))
        r += 1
        d += timedelta(days=1)
    return p_r1, p_r2, days


def test_the_pay_sheet_exists(wb):
    assert "Stipendio" in wb.sheetnames


def test_table_heading(wb):
    ws = wb["Stipendio"]
    attese = ["Mese", "Ore ord.", "Ore sab", "Ore dom", "Ore fest",
              "Lordo base", "Magg. sab", "Magg. dom", "Magg. fest",
              "Rateo 13a", "Lordo totale", "Netto stimato"]
    assert [ws.cell(row=15, column=c).value for c in range(1, 13)] == attese


def test_the_twelve_months_in_order(wb):
    ws = wb["Stipendio"]
    assert [ws.cell(row=r, column=1).value for r in range(16, 28)] == common.MONTHS


def test_december_holiday_hours(wb):
    # December 2026 has three holidays: 8, 25 and 26. N() so that a holiday
    # not worked yet (empty Presenze) counts zero hours, not #VALUE!.
    ws = wb["Stipendio"]
    righe = [common.rota_row(date(YEAR, 12, g), YEAR) for g in (8, 25, 26)]
    atteso = "=" + "+".join(f"N(Presenze!$E${r})" for r in righe)
    assert ws.cell(row=MONTH_ROWS["Dicembre"], column=5).value == atteso


def test_a_month_without_holidays_has_zero_holiday_hours(wb):
    # In 2026 July, September and October have no national holidays.
    ws = wb["Stipendio"]
    for mese in ("Luglio", "Settembre", "Ottobre"):
        assert ws.cell(row=MONTH_ROWS[mese], column=5).value == 0, mese


def test_saturday_hours_exclude_saturday_holidays(wb):
    # 25 April 2026 is a Saturday and a holiday: it must not sit in Saturday hours.
    ws = wb["Stipendio"]
    riga_25 = common.rota_row(date(YEAR, 4, 25), YEAR)
    r = MONTH_ROWS["Aprile"]
    atteso = (f'=SUMIFS(Presenze!$E$5:$E$369,Presenze!$C$5:$C$369,$A{r},'
              f'Presenze!$B$5:$B$369,"Sabato")-N(Presenze!$E${riga_25})')
    assert ws.cell(row=r, column=3).value == atteso


def test_sunday_hours_exclude_sunday_holidays(wb):
    # 1 November 2026 (Ognissanti) falls on a Sunday.
    ws = wb["Stipendio"]
    riga_1 = common.rota_row(date(YEAR, 11, 1), YEAR)
    r = MONTH_ROWS["Novembre"]
    atteso = (f'=SUMIFS(Presenze!$E$5:$E$369,Presenze!$C$5:$C$369,$A{r},'
              f'Presenze!$B$5:$B$369,"Domenica")-N(Presenze!$E${riga_1})')
    assert ws.cell(row=r, column=4).value == atteso


def test_ordinary_hours_are_the_remainder(wb):
    ws = wb["Stipendio"]
    r = MONTH_ROWS["Gennaio"]
    assert ws.cell(row=r, column=2).value == (
        f'=SUMIFS(Presenze!$E$5:$E$369,Presenze!$C$5:$C$369,$A{r})'
        f'-C{r}-D{r}-E{r}')


def test_the_check_row_watches_ordinary_hours(wb):
    # B28+C28+D28+E28-SUM(hours) is identically 0 by construction (B is the
    # difference SUMIFS(month)-C-D-E), so it is no check at all: it would
    # never notice a day counted twice. The real check is the minimum of the
    # monthly ordinary hours: a double count drains them below zero.
    ws = wb["Stipendio"]
    assert ws.cell(row=30, column=1).value == (
        "Controllo ore ordinarie (minimo mensile, deve essere >= 0)")
    assert ws.cell(row=30, column=2).value == "=MIN(B16:B27)"


def test_the_check_row_turns_red_when_negative(wb):
    ws = wb["Stipendio"]
    regole = []
    for intervallo in ws.conditional_formatting:
        if "B30" in str(intervallo.sqref).split():
            regole.extend(intervallo.rules)
    assert any(
        r.operator == "lessThan" and r.formula == ["0"] for r in regole), regole


def test_setting_cells_are_empty_and_green(wb):
    ws = wb["Stipendio"]
    etichette = [ws.cell(row=r, column=1).value for r in range(5, 11)]
    assert etichette == [
        "Tariffa oraria lorda", "Maggiorazione sabato", "Maggiorazione domenica",
        "Maggiorazione festivo", "Rateo 13a", "Coefficiente netto/lordo",
    ]
    for r in (5, 6, 7, 8, 10):
        c = ws.cell(row=r, column=2)
        assert c.value is None, r
        assert c.fill.fgColor.rgb[-6:] == "E8F3EA", r
    assert ws["B9"].value == "=1/12"


def test_amounts_stay_empty_until_the_rate_is_set(wb):
    ws = wb["Stipendio"]
    r = MONTH_ROWS["Gennaio"]
    assert ws.cell(row=r, column=6).value == (
        f'=IF($B$5="","",(B{r}+C{r}+D{r}+E{r})*$B$5)')
    assert ws.cell(row=r, column=7).value == (
        f'=IF(OR($B$5="",$B$6=""),"",C{r}*$B$5*$B$6)')
    assert ws.cell(row=r, column=8).value == (
        f'=IF(OR($B$5="",$B$7=""),"",D{r}*$B$5*$B$7)')
    assert ws.cell(row=r, column=9).value == (
        f'=IF(OR($B$5="",$B$8=""),"",E{r}*$B$5*$B$8)')
    assert ws.cell(row=r, column=10).value == (
        f'=IF(OR(G{r}="",H{r}="",I{r}="",$B$9=""),"",(F{r}+G{r}+H{r}+I{r})*$B$9)')
    assert ws.cell(row=r, column=11).value == (
        f'=IF(J{r}="","",F{r}+G{r}+H{r}+I{r}+J{r})')


def test_the_net_also_needs_the_ratio(wb):
    ws = wb["Stipendio"]
    r = MONTH_ROWS["Gennaio"]
    assert ws.cell(row=r, column=12).value == (
        f'=IF(OR(K{r}="",$B$10=""),"",K{r}*$B$10)')


def test_the_yearly_total_stays_empty_until_the_rate_is_set(wb):
    # Found in review: SUM() ignores empty-text cells and would return 0 on
    # the whole TOTALE ANNO row when the sheet is not configured yet, the
    # same misleading figure the design wants to avoid, moved from a single
    # cell to the yearly total.
    ws = wb["Stipendio"]
    for col in range(6, 13):
        L = get_column_letter(col)
        assert ws.cell(row=28, column=col).value == (
            f'=IF({L}16="","",SUM({L}16:{L}27))'), col


def test_euro_format_on_the_amounts(wb):
    ws = wb["Stipendio"]
    r = MONTH_ROWS["Gennaio"]
    for col in range(6, 13):
        assert ws.cell(row=r, column=col).number_format == '€ #,##0.00', col


def test_the_chart_is_there(wb):
    ws = wb["Stipendio"]
    assert len(ws._charts) == 1
    assert ws._charts[0].title is not None


def test_the_note_about_what_is_not_covered(wb):
    ws = wb["Stipendio"]
    testi = [ws.cell(row=r, column=1).value for r in range(1, 15)]
    assert any(isinstance(t, str) and "straordinari" in t for t in testi)
    assert any(isinstance(t, str) and "busta paga" in t for t in testi)


# --- Real recalculation with LibreOffice ---------------------------------
# The tests above compare formula strings only: that is how the yearly-total
# defect (SUM ignores empty text and returns 0) slipped through review.
# Here the numbers are actually opened.
SOFFICE = shutil.which("soffice")


def _build_workbook_for_recalc(tmp_path, name, settings, uniform_hours=None):
    """Generates a fresh workbook and writes literal values into the green
    setting cells and, on request, into each month's hours.

    Each monthly row's hours are overwritten with the same quartet of literal
    numbers (no longer formulas) on purpose: it isolates the arithmetic of
    the euro columns, which is what this test is about, from the behaviour of
    the hour formulas themselves.
    """
    out = tmp_path / f"{name}.xlsx"
    subprocess.run(
        [sys.executable, "generate_rota.py", str(YEAR), str(out)],
        cwd=ROOT, check=True, capture_output=True, timeout=60,
    )
    wb = load_workbook(out)
    ws = wb["Stipendio"]
    for cell, value in settings.items():
        ws[cell] = value
    if uniform_hours is not None:
        ordinary, sat, sun, hol = uniform_hours
        for r in range(pay_sheet.FIRST_MONTH_ROW, pay_sheet.TOTAL_ROW):
            ws.cell(row=r, column=2, value=ordinary)
            ws.cell(row=r, column=3, value=sat)
            ws.cell(row=r, column=4, value=sun)
            ws.cell(row=r, column=5, value=hol)
    wb.save(out)
    return out


def _recalculate_with_soffice(path, tmp_path):
    """Converts with headless LibreOffice (which recalculates the formulas)
    and reloads the result with data_only=True: the cells then hold the
    computed value rather than the formula."""
    converted = tmp_path / "convertiti"
    converted.mkdir(exist_ok=True)
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "xlsx",
         "--outdir", str(converted), str(path)],
        check=True, capture_output=True, timeout=120,
    )
    return load_workbook(converted / path.name, data_only=True)


@pytest.mark.skipif(
    SOFFICE is None,
    reason="LibreOffice (soffice) is not on PATH: it is needed to recalculate.")
def test_euro_values_are_correct_after_recalculation(tmp_path):
    ws_col = list(range(6, 13))  # F..L

    # (a) A wholly unconfigured sheet: months and yearly total stay blank in
    # every euro column, not zero.
    empty_path = _build_workbook_for_recalc(tmp_path, "empty", {})
    wb_empty = _recalculate_with_soffice(empty_path, tmp_path)
    ws = wb_empty["Stipendio"]
    for col in ws_col:
        assert ws.cell(row=16, column=col).value is None, (
            "mese", col, ws.cell(row=16, column=col).value)
        assert ws.cell(row=28, column=col).value is None, (
            "totale", col, ws.cell(row=28, column=col).value)

    # (b) A wholly configured sheet, with round numbers: every euro column
    # must match a calculation written out by hand here, not taken from
    # pay_sheet.py. The hours are the same every month, so the yearly total
    # is simply twelve times the monthly value.
    ordinary, sat, sun, hol = 100, 10, 5, 2
    rate, sat_pct, sun_pct, hol_pct, accrual, net_ratio = 10, 0.2, 0.3, 0.5, 0.1, 0.75
    full_path = _build_workbook_for_recalc(
        tmp_path, "full",
        {"B5": rate, "B6": sat_pct, "B7": sun_pct, "B8": hol_pct,
         "B9": accrual, "B10": net_ratio},
        uniform_hours=(ordinary, sat, sun, hol))
    wb_full = _recalculate_with_soffice(full_path, tmp_path)
    ws = wb_full["Stipendio"]

    expected_base = (ordinary + sat + sun + hol) * rate
    expected_sat = sat * rate * sat_pct
    expected_sun = sun * rate * sun_pct
    expected_hol = hol * rate * hol_pct
    expected_accrual = (expected_base + expected_sat + expected_sun + expected_hol) * accrual
    expected_gross = expected_base + expected_sat + expected_sun + expected_hol + expected_accrual
    expected_net = expected_gross * net_ratio
    expected_month = [expected_base, expected_sat, expected_sun, expected_hol,
                      expected_accrual, expected_gross, expected_net]

    for col, expected in zip(ws_col, expected_month):
        assert ws.cell(row=16, column=col).value == pytest.approx(expected), (
            "month", col)
    for col, expected in zip(ws_col, expected_month):
        assert ws.cell(row=28, column=col).value == pytest.approx(expected * 12), (
            "total", col)

    # (c) Rate set but premiums not filled in yet: those columns stay blank,
    # not zero, neither in the month row nor in the yearly total.
    partial_path = _build_workbook_for_recalc(
        tmp_path, "partial", {"B5": rate},
        uniform_hours=(ordinary, sat, sun, hol))
    wb_partial = _recalculate_with_soffice(partial_path, tmp_path)
    ws = wb_partial["Stipendio"]

    assert ws.cell(row=16, column=6).value == pytest.approx(expected_base)
    assert ws.cell(row=28, column=6).value == pytest.approx(expected_base * 12)
    for col in range(7, 13):  # G..L: premiums, accrual, gross total, net
        assert ws.cell(row=16, column=col).value is None, ("month", col)
        assert ws.cell(row=28, column=col).value is None, ("total", col)


_RE_ERROR = re.compile(r'^#[A-Z/0-9!]+$')  # #VALUE!, #DIV/0!, #N/A, etc.


@pytest.mark.skipif(
    SOFFICE is None,
    reason="LibreOffice (soffice) is not on PATH: it is needed to recalculate.")
def test_hours_are_zero_not_an_error_when_the_rota_is_empty(tmp_path):
    """The very first run: a freshly generated workbook, Presenze still
    empty, but the pay settings already filled in (that is the first thing a
    person does). The hour columns (B..E) must read as "zero hours worked",
    a true fact, not as a calculation error.
    """
    p = _build_workbook_for_recalc(
        tmp_path, "empty_rota",
        {"B5": 10, "B6": 0.2, "B7": 0.3, "B8": 0.5, "B9": 0.1, "B10": 0.75})
    # No uniform_hours here on purpose: what is wanted is the real behaviour
    # of the B..E formulas when Presenze holds no code yet.
    wb_conv = _recalculate_with_soffice(p, tmp_path)
    ws = wb_conv["Stipendio"]

    for r in range(pay_sheet.FIRST_MONTH_ROW, pay_sheet.TOTAL_ROW + 1):
        for col in range(1, 13):
            v = ws.cell(row=r, column=col).value
            assert not (isinstance(v, str) and _RE_ERROR.match(v)), (r, col, v)

    for r in range(pay_sheet.FIRST_MONTH_ROW, pay_sheet.TOTAL_ROW):
        for col in range(2, 6):  # B..E: ordinary/saturday/sunday/holiday hours
            assert ws.cell(row=r, column=col).value == 0, (r, col)

    check = ws.cell(row=pay_sheet.CHECK_ROW, column=2).value
    assert check == 0, check


def test_every_day_of_the_year_lands_in_exactly_one_bucket(wb):
    """Columns C/D/E/B (sat/sun/hol/ord) must partition each month's Presenze
    rows exactly: every day once, never twice, never zero. The expected
    bucket comes from the real calendar (italian_holidays plus the weekday
    from date.weekday()), not from the formula itself: that way a bug moving
    or duplicating a day between buckets is caught, which comparing strings
    built by the same production code could never do.
    """
    ws = wb["Stipendio"]
    p_r1, p_r2, days = _year_bounds(YEAR)
    holidays = set(common.italian_holidays(YEAR))
    hours_rng = f"Presenze!$E${p_r1}:$E${p_r2}"
    months_rng = f"Presenze!$C${p_r1}:$C${p_r2}"
    weekdays_rng = f"Presenze!$B${p_r1}:$B${p_r2}"

    re_sumifs_weekday = re.compile(
        r'^=SUMIFS\(' + re.escape(hours_rng) + ',' + re.escape(months_rng)
        + r',\$A(\d+),' + re.escape(weekdays_rng) + r',"(Sabato|Domenica)"\)'
        r'((?:-N\(Presenze!\$E\$\d+\))*)$')
    re_sumifs_month = re.compile(
        r'^=SUMIFS\(' + re.escape(hours_rng) + ',' + re.escape(months_rng)
        + r',\$A(\d+)\)-C(\d+)-D(\d+)-E(\d+)$')

    for i, name in enumerate(common.MONTHS):
        r = MONTH_ROWS[name]
        month_num = i + 1
        month_days = [(rr, d) for rr, d in days if d.month == month_num]

        # Column E: explicit sum of holiday cells, or the integer 0.
        hol_formula = ws.cell(row=r, column=5).value
        hol = {}
        if hol_formula != 0:
            assert isinstance(hol_formula, str) and hol_formula.startswith("="), (name, hol_formula)
            for term in hol_formula[1:].split("+"):
                m = _RE_HOLIDAY_TERM.match(term)
                assert m, (name, "unrecognised holiday term", term)
                hol[int(m.group(1))] = hol.get(int(m.group(1)), 0) + 1

        # Columns C/D: SUMIFS on the weekday, minus the explicit subtraction
        # of holidays falling on that weekday. The set of rows the SUMIFS
        # touches is derived from the real calendar, not from the sheet.
        buckets = {"hol": hol}
        for col, key, weekday_name in ((3, "sat", "Sabato"), (4, "sun", "Domenica")):
            v = ws.cell(row=r, column=col).value
            m = re_sumifs_weekday.match(v)
            assert m, (name, key, v)
            row_ref, expected_weekday = m.group(1), m.group(2)
            assert int(row_ref) == r and expected_weekday == weekday_name, (name, v)
            bucket = {}
            for rr, d in month_days:
                if common.WEEKDAYS[d.weekday()] == weekday_name:
                    bucket[rr] = bucket.get(rr, 0) + 1
            for n in _RE_SUBTRACTION.findall(m.group(3)):
                nn = int(n)
                bucket[nn] = bucket.get(nn, 0) - 1
            buckets[key] = bucket

        # Column B: ordinary hours as the remainder. Check the formula really
        # has the SUMIFS(month)-C-D-E shape pointing at the right row, then
        # derive its per-row contribution from what was already computed for
        # sat/sun/hol (1 if the day is in the month, less what was taken
        # elsewhere).
        ord_formula = ws.cell(row=r, column=2).value
        m = re_sumifs_month.match(ord_formula)
        assert m, (name, ord_formula)
        assert tuple(int(g) for g in m.groups()) == (r, r, r, r), (name, ord_formula)
        buckets["ord"] = {
            rr: 1 - buckets["sat"].get(rr, 0) - buckets["sun"].get(rr, 0)
            - buckets["hol"].get(rr, 0)
            for rr, _ in month_days
        }

        # Partition check: every day of the month in exactly one bucket, and
        # precisely the one the real calendar dictates (precedence
        # holiday > sunday > saturday > ordinary).
        for rr, d in month_days:
            is_holiday = d in holidays
            expected = {
                "hol": 1 if is_holiday else 0,
                "sun": 1 if (not is_holiday and d.weekday() == 6) else 0,
                "sat": 1 if (not is_holiday and d.weekday() == 5) else 0,
            }
            expected["ord"] = 1 - expected["hol"] - expected["sun"] - expected["sat"]
            for key, want in expected.items():
                got = buckets[key].get(rr, 0)
                assert got == want, (
                    name, d, key, "expected", want, "got", got)
            total = sum(buckets[k].get(rr, 0) for k in ("hol", "sun", "sat", "ord"))
            assert total == 1, (name, d, "sum of buckets", total)
