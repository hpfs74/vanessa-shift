"""Test del foglio Stipendio."""

import re
import shutil
import subprocess
import sys
from datetime import date, timedelta

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

import comune
import stipendio
from conftest import ANNO, RADICE

MESI_RIGHE = {m: 16 + i for i, m in enumerate(comune.MESI)}

# --- pattern delle sole forme di formula che stipendio.py puo' produrre nelle
# colonne B/C/D/E. Servono al test di partizione qui sotto: dato un anno,
# riconoscono qualunque intervallo Presenze!$E$p_r1:$E$p_r2 (non fisso
# 5:369), cosi' il test resta valido anche cambiando anno. ---
_RE_FEST_TERMINE = re.compile(r'^Presenze!\$E\$(\d+)$')
_RE_SOTTRAZIONE = re.compile(r'-Presenze!\$E\$(\d+)')


def _confini_anno(anno):
    """Prima e ultima riga di Presenze per l'anno, e la lista (riga, data)."""
    p_r1 = comune.riga_presenze(date(anno, 1, 1), anno)
    p_r2 = comune.riga_presenze(date(anno, 12, 31), anno)
    giorni, d, r = [], date(anno, 1, 1), p_r1
    while r <= p_r2:
        giorni.append((r, d))
        r += 1
        d += timedelta(days=1)
    return p_r1, p_r2, giorni


def test_il_foglio_stipendio_esiste(wb):
    assert "Stipendio" in wb.sheetnames


def test_intestazione_della_tabella(wb):
    ws = wb["Stipendio"]
    attese = ["Mese", "Ore ord.", "Ore sab", "Ore dom", "Ore fest",
              "Lordo base", "Magg. sab", "Magg. dom", "Magg. fest",
              "Rateo 13a", "Lordo totale", "Netto stimato"]
    assert [ws.cell(row=15, column=c).value for c in range(1, 13)] == attese


def test_i_dodici_mesi_in_ordine(wb):
    ws = wb["Stipendio"]
    assert [ws.cell(row=r, column=1).value for r in range(16, 28)] == comune.MESI


def test_ore_festive_di_dicembre(wb):
    # Dicembre 2026 ha tre festivi: 8, 25 e 26.
    ws = wb["Stipendio"]
    righe = [comune.riga_presenze(date(ANNO, 12, g), ANNO) for g in (8, 25, 26)]
    atteso = "=" + "+".join(f"Presenze!$E${r}" for r in righe)
    assert ws.cell(row=MESI_RIGHE["Dicembre"], column=5).value == atteso


def test_mese_senza_festivi_ha_zero_ore_festive(wb):
    # Nel 2026 luglio, settembre e ottobre non hanno festivi nazionali.
    ws = wb["Stipendio"]
    for mese in ("Luglio", "Settembre", "Ottobre"):
        assert ws.cell(row=MESI_RIGHE[mese], column=5).value == 0, mese


def test_le_ore_di_sabato_scartano_i_festivi_di_sabato(wb):
    # 25 aprile 2026 e' sabato ed e' festivo: non deve stare fra le ore di sabato.
    ws = wb["Stipendio"]
    riga_25 = comune.riga_presenze(date(ANNO, 4, 25), ANNO)
    r = MESI_RIGHE["Aprile"]
    atteso = (f'=SUMIFS(Presenze!$E$5:$E$369,Presenze!$C$5:$C$369,$A{r},'
              f'Presenze!$B$5:$B$369,"Sabato")-Presenze!$E${riga_25}')
    assert ws.cell(row=r, column=3).value == atteso


def test_le_ore_di_domenica_scartano_i_festivi_di_domenica(wb):
    # 1 novembre 2026 (Ognissanti) cade di domenica.
    ws = wb["Stipendio"]
    riga_1 = comune.riga_presenze(date(ANNO, 11, 1), ANNO)
    r = MESI_RIGHE["Novembre"]
    atteso = (f'=SUMIFS(Presenze!$E$5:$E$369,Presenze!$C$5:$C$369,$A{r},'
              f'Presenze!$B$5:$B$369,"Domenica")-Presenze!$E${riga_1}')
    assert ws.cell(row=r, column=4).value == atteso


def test_ore_ordinarie_sono_il_resto(wb):
    ws = wb["Stipendio"]
    r = MESI_RIGHE["Gennaio"]
    assert ws.cell(row=r, column=2).value == (
        f'=SUMIFS(Presenze!$E$5:$E$369,Presenze!$C$5:$C$369,$A{r})'
        f'-C{r}-D{r}-E{r}')


def test_riga_di_controllo_confronta_con_le_ore_totali(wb):
    # B28+C28+D28+E28-SUM(ore) e' identicamente 0 per costruzione (B e' la
    # differenza SUMIFS(mese)-C-D-E), quindi non e' un controllo: non si
    # accorgerebbe mai di un giorno contato due volte. Il controllo vero e'
    # il minimo delle ore ordinarie mensili: un doppio conteggio le drena
    # sotto zero.
    ws = wb["Stipendio"]
    assert ws.cell(row=30, column=1).value == (
        "Controllo ore ordinarie (minimo mensile, deve essere >= 0)")
    assert ws.cell(row=30, column=2).value == "=MIN(B16:B27)"


def test_la_riga_di_controllo_si_colora_se_negativa(wb):
    ws = wb["Stipendio"]
    regole = []
    for intervallo in ws.conditional_formatting:
        if "B30" in str(intervallo.sqref).split():
            regole.extend(intervallo.rules)
    assert any(
        r.operator == "lessThan" and r.formula == ["0"] for r in regole), regole


def test_celle_dei_parametri_sono_vuote_e_verdi(wb):
    ws = wb["Stipendio"]
    etichette = [ws.cell(row=r, column=1).value for r in range(5, 11)]
    assert etichette == [
        "Tariffa oraria lorda", "Maggiorazione sabato", "Maggiorazione domenica",
        "Maggiorazione festivo", "Rateo 13a", "Coefficiente netto/lordo",
    ]
    for r in (5, 6, 7, 8, 10):
        c = ws.cell(row=r, column=2)
        assert c.value is None, r
        assert c.fill.fgColor.rgb[-6:] == "E8F3EA", r
    assert ws["B9"].value == "=1/12"


def test_gli_importi_restano_vuoti_finche_manca_la_tariffa(wb):
    ws = wb["Stipendio"]
    r = MESI_RIGHE["Gennaio"]
    assert ws.cell(row=r, column=6).value == (
        f'=IF($B$5="","",(B{r}+C{r}+D{r}+E{r})*$B$5)')
    assert ws.cell(row=r, column=7).value == (
        f'=IF(OR($B$5="",$B$6=""),"",C{r}*$B$5*$B$6)')
    assert ws.cell(row=r, column=8).value == (
        f'=IF(OR($B$5="",$B$7=""),"",D{r}*$B$5*$B$7)')
    assert ws.cell(row=r, column=9).value == (
        f'=IF(OR($B$5="",$B$8=""),"",E{r}*$B$5*$B$8)')
    assert ws.cell(row=r, column=10).value == (
        f'=IF(OR(G{r}="",H{r}="",I{r}="",$B$9=""),"",(F{r}+G{r}+H{r}+I{r})*$B$9)')
    assert ws.cell(row=r, column=11).value == (
        f'=IF(J{r}="","",F{r}+G{r}+H{r}+I{r}+J{r})')


def test_il_netto_richiede_anche_il_coefficiente(wb):
    ws = wb["Stipendio"]
    r = MESI_RIGHE["Gennaio"]
    assert ws.cell(row=r, column=12).value == (
        f'=IF(OR(K{r}="",$B$10=""),"",K{r}*$B$10)')


def test_il_totale_anno_resta_vuoto_finche_manca_la_tariffa(wb):
    # Trovato in revisione: SUM() ignora le celle di testo "" e restituirebbe
    # 0 sull'intera riga TOTALE ANNO quando il foglio non e' ancora
    # configurato, cioe' la stessa cifra fuorviante che il design vuole
    # evitare, spostata dalla singola cella al totale annuo.
    ws = wb["Stipendio"]
    for col in range(6, 13):
        L = get_column_letter(col)
        assert ws.cell(row=28, column=col).value == (
            f'=IF({L}16="","",SUM({L}16:{L}27))'), col


def test_formato_euro_sugli_importi(wb):
    ws = wb["Stipendio"]
    r = MESI_RIGHE["Gennaio"]
    for col in range(6, 13):
        assert ws.cell(row=r, column=col).number_format == '€ #,##0.00', col


def test_ce_il_grafico(wb):
    ws = wb["Stipendio"]
    assert len(ws._charts) == 1
    assert ws._charts[0].title is not None


def test_nota_sui_limiti(wb):
    ws = wb["Stipendio"]
    testi = [ws.cell(row=r, column=1).value for r in range(1, 15)]
    assert any(isinstance(t, str) and "straordinari" in t for t in testi)
    assert any(isinstance(t, str) and "busta paga" in t for t in testi)


# --- Ricalcolo vero con LibreOffice --------------------------------------
# I test sopra confrontano solo le stringhe delle formule: e' cosi' che il
# difetto del totale annuo (SUM ignora il testo vuoto e restituisce 0) e'
# passato inosservato in revisione. Qui si aprono davvero i numeri.
SOFFICE = shutil.which("soffice")


def _genera_stipendio_per_ricalcolo(tmp_path, nome, parametri, ore_uniformi=None):
    """Genera un workbook fresco e vi scrive valori letterali nelle celle
    verdi dei parametri e, se richiesto, nelle ore di ogni mese.

    Le ore di ogni riga mensile vengono sovrascritte con lo stesso quartetto
    di numeri letterali (non piu' formule) apposta: isola l'aritmetica delle
    colonne in euro, oggetto di questo test, da un problema distinto e
    preesistente non toccato in questo giro di correzioni, per cui a foglio
    Presenze vuoto la colonna "Ore festive" (somma diretta di celle
    Presenze!$E$n che sono ancora testo vuoto) produce #VALUE!.
    """
    out = tmp_path / f"{nome}.xlsx"
    subprocess.run(
        [sys.executable, "genera_presenze.py", str(ANNO), str(out)],
        cwd=RADICE, check=True, capture_output=True, timeout=60,
    )
    wb = load_workbook(out)
    ws = wb["Stipendio"]
    for cella, valore in parametri.items():
        ws[cella] = valore
    if ore_uniformi is not None:
        ordinarie, sab, dom, fest = ore_uniformi
        for r in range(stipendio.RIGA_PRIMO_MESE, stipendio.RIGA_TOTALE):
            ws.cell(row=r, column=2, value=ordinarie)
            ws.cell(row=r, column=3, value=sab)
            ws.cell(row=r, column=4, value=dom)
            ws.cell(row=r, column=5, value=fest)
    wb.save(out)
    return out


def _ricalcola_con_soffice(path, tmp_path):
    """Converte con LibreOffice headless (che ricalcola le formule) e
    ricarica il risultato con data_only=True: le celle hanno il valore
    calcolato, non piu' la formula."""
    convertiti = tmp_path / "convertiti"
    convertiti.mkdir(exist_ok=True)
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "xlsx",
         "--outdir", str(convertiti), str(path)],
        check=True, capture_output=True, timeout=120,
    )
    return load_workbook(convertiti / path.name, data_only=True)


@pytest.mark.skipif(
    SOFFICE is None,
    reason="LibreOffice (soffice) non e' sul PATH: serve a ricalcolare le formule.")
def test_i_valori_in_euro_sono_corretti_dopo_il_ricalcolo(tmp_path):
    ws_col = list(range(6, 13))  # F..L

    # (a) Foglio del tutto non configurato: mesi e totale annuo restano
    # vuoti in ogni colonna in euro, non zero.
    p_vuoto = _genera_stipendio_per_ricalcolo(tmp_path, "vuoto", {})
    wb_vuoto = _ricalcola_con_soffice(p_vuoto, tmp_path)
    ws = wb_vuoto["Stipendio"]
    for col in ws_col:
        assert ws.cell(row=16, column=col).value is None, (
            "mese", col, ws.cell(row=16, column=col).value)
        assert ws.cell(row=28, column=col).value is None, (
            "totale", col, ws.cell(row=28, column=col).value)

    # (b) Foglio del tutto configurato, con numeri tondi: ogni colonna in
    # euro deve corrispondere a un calcolo scritto qui a mano, non preso da
    # stipendio.py. Le ore sono le stesse per ogni mese, cosi' il totale
    # annuo e' semplicemente dodici volte il valore mensile.
    ordinarie, sab, dom, fest = 100, 10, 5, 2
    tariffa, magg_sab, magg_dom, magg_fest, rateo, coeff = 10, 0.2, 0.3, 0.5, 0.1, 0.75
    p_pieno = _genera_stipendio_per_ricalcolo(
        tmp_path, "pieno",
        {"B5": tariffa, "B6": magg_sab, "B7": magg_dom, "B8": magg_fest,
         "B9": rateo, "B10": coeff},
        ore_uniformi=(ordinarie, sab, dom, fest))
    wb_pieno = _ricalcola_con_soffice(p_pieno, tmp_path)
    ws = wb_pieno["Stipendio"]

    base_attesa = (ordinarie + sab + dom + fest) * tariffa
    sab_attesa = sab * tariffa * magg_sab
    dom_attesa = dom * tariffa * magg_dom
    fest_attesa = fest * tariffa * magg_fest
    rateo_atteso = (base_attesa + sab_attesa + dom_attesa + fest_attesa) * rateo
    lordo_atteso = base_attesa + sab_attesa + dom_attesa + fest_attesa + rateo_atteso
    netto_atteso = lordo_atteso * coeff
    attesi_mese = [base_attesa, sab_attesa, dom_attesa, fest_attesa,
                   rateo_atteso, lordo_atteso, netto_atteso]

    for col, atteso in zip(ws_col, attesi_mese):
        assert ws.cell(row=16, column=col).value == pytest.approx(atteso), (
            "mese", col)
    for col, atteso in zip(ws_col, attesi_mese):
        assert ws.cell(row=28, column=col).value == pytest.approx(atteso * 12), (
            "totale", col)

    # (c) Tariffa impostata ma maggiorazioni non ancora compilate: quelle
    # colonne restano vuote, non zero - ne' nella riga del mese ne' nel
    # totale annuo.
    p_parziale = _genera_stipendio_per_ricalcolo(
        tmp_path, "parziale", {"B5": tariffa},
        ore_uniformi=(ordinarie, sab, dom, fest))
    wb_parziale = _ricalcola_con_soffice(p_parziale, tmp_path)
    ws = wb_parziale["Stipendio"]

    assert ws.cell(row=16, column=6).value == pytest.approx(base_attesa)
    assert ws.cell(row=28, column=6).value == pytest.approx(base_attesa * 12)
    for col in range(7, 13):  # G..L: maggiorazioni, rateo, lordo totale, netto
        assert ws.cell(row=16, column=col).value is None, ("mese", col)
        assert ws.cell(row=28, column=col).value is None, ("totale", col)


def test_ogni_giorno_dell_anno_sta_in_esattamente_un_secchio(wb):
    """Le colonne C/D/E/B (sab/dom/fest/ord) devono partizionare esattamente
    le righe di Presenze di ogni mese: ogni giorno una volta sola, mai due,
    mai zero. La categoria attesa viene dal calendario vero
    (festivi_italiani + giorno della settimana calcolato da date.weekday()),
    non dalla formula stessa: cosi' un bug che sposta o duplica un giorno fra
    i secchi viene scoperto, cosa che confrontare stringhe costruite con lo
    stesso codice di produzione non potrebbe mai fare.
    """
    ws = wb["Stipendio"]
    p_r1, p_r2, giorni = _confini_anno(ANNO)
    festivi = set(comune.festivi_italiani(ANNO))
    ore_rng = f"Presenze!$E${p_r1}:$E${p_r2}"
    mesi_rng = f"Presenze!$C${p_r1}:$C${p_r2}"
    giorni_rng = f"Presenze!$B${p_r1}:$B${p_r2}"

    re_sumifs_giorno = re.compile(
        r'^=SUMIFS\(' + re.escape(ore_rng) + ',' + re.escape(mesi_rng)
        + r',\$A(\d+),' + re.escape(giorni_rng) + r',"(Sabato|Domenica)"\)'
        r'((?:-Presenze!\$E\$\d+)*)$')
    re_sumifs_mese = re.compile(
        r'^=SUMIFS\(' + re.escape(ore_rng) + ',' + re.escape(mesi_rng)
        + r',\$A(\d+)\)-C(\d+)-D(\d+)-E(\d+)$')

    for i, nome in enumerate(comune.MESI):
        r = MESI_RIGHE[nome]
        mese_num = i + 1
        giorni_mese = [(rr, d) for rr, d in giorni if d.month == mese_num]

        # Colonna E: somma esplicita di celle festivo, o l'intero 0.
        v_fest = ws.cell(row=r, column=5).value
        fest = {}
        if v_fest != 0:
            assert isinstance(v_fest, str) and v_fest.startswith("="), (nome, v_fest)
            for termine in v_fest[1:].split("+"):
                m = _RE_FEST_TERMINE.match(termine)
                assert m, (nome, "termine festivo non riconosciuto", termine)
                fest[int(m.group(1))] = fest.get(int(m.group(1)), 0) + 1

        # Colonne C/D: SUMIFS sul giorno della settimana, meno le sottrazioni
        # esplicite dei festivi che cadono su quel giorno. L'insieme di righe
        # colpite dal SUMIFS lo calcolo dal calendario vero, non dal foglio.
        secchi = {"fest": fest}
        for col, chiave, nome_giorno in ((3, "sab", "Sabato"), (4, "dom", "Domenica")):
            v = ws.cell(row=r, column=col).value
            m = re_sumifs_giorno.match(v)
            assert m, (nome, chiave, v)
            riga_a, giorno_atteso = m.group(1), m.group(2)
            assert int(riga_a) == r and giorno_atteso == nome_giorno, (nome, v)
            secchio = {}
            for rr, d in giorni_mese:
                if comune.GIORNI[d.weekday()] == nome_giorno:
                    secchio[rr] = secchio.get(rr, 0) + 1
            for n in _RE_SOTTRAZIONE.findall(m.group(3)):
                nn = int(n)
                secchio[nn] = secchio.get(nn, 0) - 1
            secchi[chiave] = secchio

        # Colonna B: ordinarie come resto. Verifico che la formula abbia
        # davvero la forma SUMIFS(mese)-C-D-E riferita alla riga giusta, poi
        # ne derivo il contributo per riga da quello gia' calcolato per
        # sab/dom/fest (1 se e' del mese, meno quanto gia' tolto altrove).
        v_ord = ws.cell(row=r, column=2).value
        m = re_sumifs_mese.match(v_ord)
        assert m, (nome, v_ord)
        assert tuple(int(g) for g in m.groups()) == (r, r, r, r), (nome, v_ord)
        secchi["ord"] = {
            rr: 1 - secchi["sab"].get(rr, 0) - secchi["dom"].get(rr, 0)
            - secchi["fest"].get(rr, 0)
            for rr, _ in giorni_mese
        }

        # Verifica di partizione: ogni giorno del mese in esattamente un
        # secchio, e proprio quello dettato dal calendario vero (precedenza
        # festivo > domenica > sabato > ordinario).
        for rr, d in giorni_mese:
            e_festivo = d in festivi
            attesi = {
                "fest": 1 if e_festivo else 0,
                "dom": 1 if (not e_festivo and d.weekday() == 6) else 0,
                "sab": 1 if (not e_festivo and d.weekday() == 5) else 0,
            }
            attesi["ord"] = 1 - attesi["fest"] - attesi["dom"] - attesi["sab"]
            for chiave, atteso in attesi.items():
                ottenuto = secchi[chiave].get(rr, 0)
                assert ottenuto == atteso, (
                    nome, d, chiave, "atteso", atteso, "ottenuto", ottenuto)
            totale = sum(secchi[k].get(rr, 0) for k in ("fest", "dom", "sab", "ord"))
            assert totale == 1, (nome, d, "somma dei secchi", totale)
